"""
helpers.py — Shared shape/text helpers for slide archetypes
===========================================================

Low-level composition primitives used by every archetype renderer: rects,
text, bullets, chips, stats, arrows, pictures with focal-aware cropping, and
styled tables. All coordinates are inches on the 13.333 x 7.5 canvas.
"""

from __future__ import annotations

from typing import Optional

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from mini_kio.core.presentation.design import lighter, px


def _mk_shape(slide, shape_type, x, y, w, h):
    return slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))


def fill_shape(shape, color, line_color=None, line_w: float = 0.75):
    if color is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = px(color)
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = px(line_color)
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def add_rect(slide, x, y, w, h, fill, line_color=None, line_w=0.75,
             shape_type=MSO_SHAPE.RECTANGLE, radius=0.12):
    shp = _mk_shape(slide, shape_type, x, y, w, h)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:  # noqa: BLE001
            pass
    fill_shape(shp, fill, line_color, line_w)
    shp.text_frame.word_wrap = True
    return shp


def add_text(slide, x, y, w, h, text, *, font, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, word_wrap=True,
             space_after=0.0, line_spacing=1.0, italic=False, font_color=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    if line_spacing != 1.0:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = str(text or "")
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = px(font_color or color)
    return tb


def add_multi_bullets(slide, x, y, w, h, items, *, font, size, color, marker_color,
                      bullet_char="\u2022", space_after=5.0, line_spacing=1.08,
                      max_items: Optional[int] = None):
    """A body of bullets with colored markers."""
    if max_items and len(items) > max_items:
        items = items[:max_items]
    # Truncate individual bullets that are too long to prevent text overflow
    _max_bullet_chars = 80
    items = [item[:_max_bullet_chars] + ("..." if len(item) > _max_bullet_chars else "") for item in items]
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        m = p.add_run()
        m.text = f"{bullet_char}  "
        m.font.name = font
        m.font.size = Pt(size)
        m.font.bold = True
        m.font.color.rgb = px(marker_color)
        r = p.add_run()
        r.text = item
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = px(color)
    return tb


def add_chip(slide, x, y, w, h, text, *, fill, text_color, font, size=10.5,
             bold=True, align=PP_ALIGN.CENTER, radius=0.5):
    shp = _mk_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = radius
    except Exception:  # noqa: BLE001
        pass
    fill_shape(shp, fill)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = str(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = px(text_color)
    return shp


def add_stat(slide, x, y, w, h, value, label, *, deck, align=PP_ALIGN.LEFT,
             value_size=None, value_color=None, label_color=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = str(value or "")
    r.font.name = deck.display_font
    r.font.size = Pt(value_size or deck.scale.stat)
    r.font.bold = True
    r.font.color.rgb = px(value_color or deck.palette["primary"])
    p2 = tf.add_paragraph()
    p2.alignment = align
    p2.space_before = Pt(2)
    r2 = p2.add_run()
    r2.text = str(label or "")
    r2.font.name = deck.body_font
    r2.font.size = Pt(deck.scale.caption + 1.0)
    r2.font.color.rgb = px(label_color or deck.palette["muted"])
    return tb


def add_arrow(slide, x1, y1, x2, y2, color, width_pt: float = 1.6):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = px(color)
    conn.line.width = Pt(width_pt)
    ln = conn.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


def add_line(slide, x1, y1, x2, y2, color, width_pt: float = 1.2, dash: str = ""):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = px(color)
    conn.line.width = Pt(width_pt)
    if dash:
        ln = conn.line._get_or_add_ln()
        from pptx.oxml.ns import qn
        d = ln.makeelement(qn("a:prstDash"), {"val": dash})
        ln.append(d)
    return conn


def add_picture_fit(slide, image: dict, x, y, w, h, *, focus=(0.5, 0.35),
                    corner: Optional[bool] = None, deck=None):
    """Place an image into a target box with focal-point-aware cropping."""
    path = (image or {}).get("path")
    if not path:
        return None
    try:
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    except Exception:  # noqa: BLE001
        return None
    src_w = float(image.get("width") or 0)
    src_h = float(image.get("height") or 0)
    if src_w > 0 and src_h > 0:
        target_aspect = w / h if h > 0 else 1.0
        src_aspect = src_w / src_h
        fx, fy = focus
        if src_aspect > target_aspect:
            crop_total = 1.0 - target_aspect / src_aspect
            crop_total = max(0.0, min(crop_total, 0.85))
            pic.crop_left = crop_total * fx
            pic.crop_right = crop_total * (1 - fx)
        else:
            crop_total = 1.0 - src_aspect / target_aspect
            crop_total = max(0.0, min(crop_total, 0.85))
            pic.crop_top = crop_total * fy
            pic.crop_bottom = crop_total * (1 - fy)
    if corner:
        try:
            geom = pic._element.spPr.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
            if geom is not None:
                geom.set("prst", "roundRect")
        except Exception:  # noqa: BLE001
            pass
    return pic


def add_table(slide, rows, x, y, w, *, deck, header_fill=None, col_widths=None,
              font_size=None, header_font_size=None, row_h=None):
    """Styled PPT table; returns the GraphicFrame. Rows: list[list[str]]."""
    if not rows:
        return None
    ncols = max(len(r) for r in rows)
    nrows = len(rows)
    row_h = row_h or 0.42
    total_h = row_h * nrows
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(total_h))
    table = gf.table
    if col_widths:
        for ci, cw in enumerate(col_widths):
            if ci < ncols:
                table.columns[ci].width = Inches(cw)
    header_fill = header_fill or deck.palette["primary"]
    fs = font_size or deck.scale.body
    hfs = header_font_size or deck.scale.body
    for ri, row in enumerate(rows):
        try:
            table.rows[ri].height = Inches(row_h)
        except Exception:  # noqa: BLE001
            pass
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            val = row[ci] if ci < len(row) else ""
            cell.text = ""
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.08)
            tf.margin_right = Inches(0.06)
            tf.margin_top = Inches(0.03)
            tf.margin_bottom = Inches(0.03)
            p = tf.paragraphs[0]
            if ri == 0:
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = px(header_fill)
            elif ri % 2 == 0:
                cell.fill.fore_color.rgb = px(lighter(deck.palette["bg"], 0.35))
            else:
                cell.fill.fore_color.rgb = px(deck.palette["surface"])
            r = p.add_run()
            r.text = str(val)
            r.font.name = deck.body_font
            if ri == 0:
                r.font.size = Pt(hfs)
                r.font.bold = True
                r.font.color.rgb = px(deck.palette["on_primary"])
            else:
                r.font.size = Pt(fs)
                r.font.color.rgb = px(deck.palette["ink"])
    return gf
