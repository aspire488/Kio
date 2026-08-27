"""
System-level gap-closure tests (2026-08-12).

Covers:
  A. deterministic close semantics — APP vs WINDOW vs BROWSER vs TAB
  B. generic CLOSE / SCOPE=ALL_APPLICATIONS (classifier + executor)
  C. desktop-action capability classes (TYPE / KEY_PRESS / SCROLL / CLICK)
     with the bounded anti-leak guards
  D. lock/unlock single authoritative system-session model
  E. "shut down the computer" stays a SYSTEM action, never an app close
  F. MCP startup is non-blocking (optional provider must not block the
     deterministic local command path)
"""

import os
import sys
import time
import unittest
from unittest import mock

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
os.environ.setdefault("KIO_TEST_MODE", "1")

from mini_kio.core.pipeline import Pipeline
from mini_kio.core.command_parser import parse_command


def _classify(text):
    p = Pipeline()
    return p._classifier.classify(text, text)


class CloseSemanticsTest(unittest.TestCase):
    """Section 6: a request to close an APPLICATION must never silently
    become a browser-tab operation; a web app closes at TAB scope."""

    def test_close_registered_native_app_is_app_scope(self):
        d = _classify("close spotify")
        self.assertEqual(d.intent_type.value, "desktop_close")
        self.assertEqual(d.action, "close_app")
        self.assertEqual(d.target, "spotify")

    def test_close_generic_app_is_app_scope(self):
        d = _classify("close notepad")
        self.assertEqual(d.action, "close_app")
        self.assertEqual(d.target, "notepad")

    def test_close_browser_is_app_scope(self):
        d = _classify("close chrome")
        self.assertEqual(d.action, "close_app")
        self.assertEqual(d.target, "chrome")

    def test_close_webapp_without_native_is_tab_scope(self):
        # chatgpt has no registered native install -> tab scope, never the
        # host browser process, never a fake native close.
        d = _classify("close chatgpt")
        self.assertEqual(d.intent_type.value, "browser_focus")
        self.assertEqual(d.action, "close_tab")

    def test_close_tab_phrasing_stays_tab_scope(self):
        d = _classify("close the chatgpt tab")
        self.assertEqual(d.action, "close_tab")

    def test_registered_webapp_with_native_identity_is_app_scope(self):
        # telegram is a registered native app -> close at application scope
        # (close_app resolves native-vs-web truthfully).
        d = _classify("close telegram")
        self.assertEqual(d.action, "close_app")


class CloseAllTest(unittest.TestCase):
    """Section 7: generic CLOSE / SCOPE=ALL_APPLICATIONS — natural wording
    converges on one canonical representation."""

    def test_close_all_apps(self):
        d = _classify("close all apps")
        self.assertEqual(d.intent_type.value, "desktop_close")
        self.assertEqual(d.action, "close_all_apps")

    def test_close_everything(self):
        d = _classify("close everything")
        self.assertEqual(d.action, "close_all_apps")

    def test_quit_all_applications(self):
        d = _classify("quit all applications")
        self.assertEqual(d.action, "close_all_apps")

    def test_kill_all_programs(self):
        d = _classify("kill all programs")
        self.assertEqual(d.action, "close_all_apps")

    def test_close_all_open_windows(self):
        d = _classify("close all open windows")
        self.assertEqual(d.action, "close_all_apps")

    def test_close_all(self):
        d = _classify("close all")
        self.assertEqual(d.action, "close_all_apps")

    def test_close_all_tabs_is_not_app_scope(self):
        # tabs are browser scope — must never converge on ALL_APPLICATIONS.
        d = _classify("close all tabs")
        self.assertNotEqual(d.action, "close_all_apps")

    def test_executor_skips_system_shells_and_reports_remaining(self):
        from mini_kio.core import app_operator

        fake_windows = [
            {"app": "Notepad", "base": "notepad", "pid": 101, "title": "Untitled"},
            {"app": "Spotify", "base": "spotify", "pid": 202, "title": "Spotify"},
            {"app": "File Explorer", "base": "explorer", "pid": 303, "title": ""},
            {"app": "Calculator", "base": "calculator", "pid": 404, "title": "Calculator"},
        ]
        closed = []
        with mock.patch(
            "mini_kio.core.desktop_state.observe_native_windows",
            side_effect=[(list(fake_windows), True), (list(fake_windows), True)],
        ), mock.patch.object(
            app_operator, "close_app",
            side_effect=lambda name, pid=None: closed.append((name, pid)) or {"success": True, "message": f"Closed {name}."},
        ):
            result = app_operator.close_all_user_apps()
        self.assertTrue(result["success"])
        closed_names = {n for n, _ in closed}
        # explorer (system shell) must never be targeted; real user apps are.
        self.assertNotIn("File Explorer", closed_names)
        self.assertIn("Notepad", closed_names)
        self.assertIn("Spotify", closed_names)
        self.assertIn("Calculator", closed_names)

    def test_executor_reports_truthful_remaining(self):
        from mini_kio.core import app_operator

        before = [
            {"app": "Notepad", "base": "notepad", "pid": 101, "title": ""},
            {"app": "Spotify", "base": "spotify", "pid": 202, "title": ""},
        ]
        after = [  # Spotify refused to close -> remains open
            {"app": "Spotify", "base": "spotify", "pid": 202, "title": ""},
        ]
        def _close(name, pid=None):
            if name == "Spotify":
                return {"success": False, "message": "Couldn't close Spotify."}
            return {"success": True, "message": f"Closed {name}."}

        with mock.patch(
            "mini_kio.core.desktop_state.observe_native_windows",
            side_effect=[(list(before), True), (list(after), True)],
        ), mock.patch.object(app_operator, "close_app", side_effect=_close):
            result = app_operator.close_all_user_apps()
        self.assertIn("Still open", result["message"])
        self.assertIn("Notepad", result["closed"])
        self.assertTrue(any("Spotify" in r for r in result["remaining"]))
        self.assertIn("Spotify", result["failed"])
        # Never claim everything closed when something remains.
        self.assertNotIn("Spotify", result["closed"])

    def test_executor_never_targets_explorer_shell(self):
        # User directive (2026-08-12): close-all must NEVER target explorer.exe
        # — File Explorer is the Windows desktop shell. Even when a File
        # Explorer window (with a real titled window, as in practice) is
        # observed, it must be excluded from the close pass AND from the
        # remaining report.
        from mini_kio.core import app_operator

        fake_windows = [
            {"app": "File Explorer", "base": "explorer", "pid": 303, "title": "Documents"},
            {"app": "Notepad", "base": "notepad", "pid": 101, "title": "Untitled"},
        ]
        closed = []
        with mock.patch(
            "mini_kio.core.desktop_state.observe_native_windows",
            side_effect=[(list(fake_windows), True), (list(fake_windows), True)],
        ), mock.patch.object(
            app_operator, "close_app",
            side_effect=lambda name, pid=None: closed.append((name, pid)) or {"success": True, "message": f"Closed {name}."},
        ):
            result = app_operator.close_all_user_apps()
        closed_names = {n for n, _ in closed}
        self.assertNotIn("File Explorer", closed_names)
        self.assertNotIn("explorer", closed_names)
        self.assertNotIn("explorer.exe", closed_names)
        self.assertNotIn("File Explorer", result.get("closed", []))
        self.assertIn("Notepad", closed_names)

    def test_executor_skips_explorer_regardless_of_base_case(self):
        # Explorer protection must be case-proof: the observation layer
        # lowercases bases, but a defensive explicit guard should hold even
        # for a mixed-case "Explorer" base.
        from mini_kio.core import app_operator

        fake_windows = [
            {"app": "File Explorer", "base": "Explorer", "pid": 303, "title": "Downloads"},
            {"app": "Calculator", "base": "calculator", "pid": 404, "title": "Calculator"},
        ]
        closed = []
        with mock.patch(
            "mini_kio.core.desktop_state.observe_native_windows",
            side_effect=[(list(fake_windows), True), (list(fake_windows), True)],
        ), mock.patch.object(
            app_operator, "close_app",
            side_effect=lambda name, pid=None: closed.append((name, pid)) or {"success": True, "message": f"Closed {name}."},
        ):
            result = app_operator.close_all_user_apps()
        closed_names = {n for n, _ in closed}
        self.assertNotIn("File Explorer", closed_names)
        self.assertNotIn("explorer", closed_names)
        self.assertIn("Calculator", closed_names)
        self.assertTrue(result["success"])

    def test_close_all_forbidden_set_is_case_normalized(self):
        # Live-found: close-all closed the terminal because the forbidden set
        # held mixed-case "WindowsTerminal.exe" while window observation
        # lowercases to "windowsterminal.exe" — the case-sensitive lookup
        # missed. Every forbidden entry must be lowercase so the protection
        # is deterministic.
        from mini_kio.core import app_operator

        forbidden = app_operator._close_all_forbidden_exes()
        self.assertTrue(all(f == f.lower() for f in forbidden))
        self.assertIn("explorer.exe", forbidden)
        self.assertIn("windowsterminal.exe", forbidden)
        self.assertIn("cmd.exe", forbidden)
        self.assertIn("dwm.exe", forbidden)

    def test_execute_action_close_all_dispatch_without_target_arg(self):
        # Live-found regression: "close all apps" failed with
        # "close_all_user_apps() takes 0 positional arguments but 1 was given"
        # because the dispatch passed the positional target. The no-arg
        # handler must be invoked as handler(), and the boundary must surface
        # a truthful failure (never a success-shaped "Done.").
        #
        # NOTE: _load_handler() returns STATIC_ACTION_TABLE[...]["handler"],
        # which captured the ORIGINAL function at import time — patching the
        # module attribute (execution_boundary.close_all_user_apps) would be
        # ineffective and the REAL close-all would run. Patch the table entry
        # the dispatcher actually resolves.
        from mini_kio.core import execution_boundary

        entry = execution_boundary.STATIC_ACTION_TABLE["close_all_apps"]
        handler = mock.Mock(return_value={"success": True, "message": "Closed all apps."})
        # execute_action builds handler_name via handler.__module__/.__name__
        handler.__name__ = "close_all_user_apps"
        handler.__module__ = "mini_kio.core.app_operator"
        with mock.patch.dict(entry, {"handler": handler}), mock.patch.object(
            execution_boundary, "_in_test_mode", return_value=False,
        ):
            result = execution_boundary.execute_action("close_all_apps", "")
        self.assertTrue(result["success"])
        handler.assert_called_once_with()  # no positional target

    def test_generic_failure_never_formats_as_done(self):
        # Live-found regression: a failed generic action was rendered as
        # "Done." — an untruthful success claim. Failures must surface the
        # actual error text.
        from mini_kio.core import runtime_response_formatter as fmt
        rendered = fmt.format_result(
            "close_all_apps", "", False,
            {"message": "Execution failed: boom", "outcome_class": "failure"},
        )
        self.assertNotIn("Done", rendered)
        self.assertIn("boom", rendered.lower())


class DesktopActionClassificationTest(unittest.TestCase):
    """Section 11/12: generic capability classes with bounded anti-leak
    guards — knowledge phrases that merely share a verb stay on the
    knowledge path."""

    def test_type_into_target(self):
        d = _classify("type hello into notepad")
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "type")
        self.assertEqual(d.metadata.get("payload"), "hello")

    def test_bare_type_short_payload(self):
        d = _classify("type hello")
        self.assertEqual(d.action, "type")
        self.assertEqual(d.metadata.get("payload"), "hello")

    def test_bare_type_long_payload_is_typing_not_conversation(self):
        # Live-found regression: "type hello world this is a kio test" leaked
        # to the conversation path, where the LLM FABRICATED a success claim
        # ("I've typed ...") instead of executing. Long typing payloads must
        # route to the canonical TYPE executor, never to conversation.
        d = _classify("type hello world this is a kio test")
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "type")
        self.assertEqual(d.metadata.get("payload"), "hello world this is a kio test")

    def test_bare_type_multiline_payload_is_typing(self):
        d = _classify("type my name is joel and i live in kochi")
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "type")

    def test_type_2_diabetes_is_knowledge_not_typing(self):
        d = _classify("type 2 diabetes")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_type_of_cancer_is_knowledge(self):
        d = _classify("type of cancer")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_type_b_blood_is_knowledge(self):
        d = _classify("type b blood")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_type_coercion_is_knowledge(self):
        d = _classify("type coercion")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_write_a_poem_is_not_typing(self):
        d = _classify("write a poem")
        # "write a poem" is CONTENT creation — never literal typing. The
        # modern classifier routes it to create_document (artifact=poem);
        # the old assertion (intent != desktop_action) is stale because
        # create_document IS a desktop_action and it is exactly the correct
        # behavior. Assert the ACTION, which encodes the actual intent.
        self.assertEqual(d.action, "create_document")
        self.assertEqual((d.metadata or {}).get("artifact"), "poem")

    def test_type_into_running_app_launches_fresh_instance_by_default(self):
        # Capability-quality (live directive): typing ALWAYS goes into a NEW
        # document — KIO never tampers with existing file content. "type hello
        # into notepad" with Notepad already running LAUNCHES a fresh app
        # instance (a new process = a guaranteed new blank document) rather
        # than typing into the existing window.
        from mini_kio.core.pipeline import _ExecutionCoordinator

        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or
            ({"success": True, "message": "ok"} if cap.startswith("keyboard_") else {"success": False, "message": "no"})
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider",
            return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus",
            return_value={"success": True, "message": "Focused Notepad."},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text", return_value=[],
        ), mock.patch.object(
            coord, "_wait_for_blank_new_window", return_value=9001,
        ), mock.patch.object(
            coord, "_read_window_text_for", return_value="",
        ), mock.patch(
            "mini_kio.core.execution_boundary.execute_action",
            return_value={"success": True, "message": "Opened notepad", "pid": 4242},
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": "hello"}},
                None,
            )
        self.assertTrue(result["success"])
        caps = [c for c, _ in events]
        # a FRESH blank window is located (never a ctrl+n hack) then typed into.
        self.assertIn("keyboard_type", caps)
        self.assertNotIn("keyboard_hotkey", caps)

    def test_type_write_onto_existing_file_reuses_window(self):
        # The only exception: an EXPLICIT write-onto request ("append",
        # "write onto the existing file") reuses the existing window — never
        # launches a second instance.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        from mini_kio.core.pipeline import RoutingDecision
        from mini_kio.core.pipeline import IntentType

        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or
            ({"success": True, "message": "ok"} if cap.startswith("keyboard_") else {"success": False, "message": "no"})
        )
        coord = _ExecutionCoordinator()
        decision = RoutingDecision(
            IntentType.DESKTOP_ACTION, "type", "notepad",
            "append this to the existing file in notepad", "append this to the existing file in notepad",
            confidence=1.0,
        )
        launched = []
        with mock.patch(
            "mini_kio.desktop.DesktopProvider",
            return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus",
            return_value={"success": True, "message": "Focused Notepad."},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text", return_value=[],
        ), mock.patch(
            "mini_kio.core.execution_boundary.execute_action",
            side_effect=lambda a, t: launched.append(a) or {"success": True, "message": "ok"},
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": "more text"}},
                decision,
            )
        self.assertTrue(result["success"])
        # write-onto: existing window reused — NO fresh launch, NO ctrl+n.
        self.assertEqual(launched, [])
        caps = [c for c, _ in events]
        self.assertNotIn("keyboard_hotkey", caps)
        self.assertIn("keyboard_type", caps)

    def test_type_new_instance_launches_fresh_process(self):
        # Explicit new-instance semantics ("type X into a NEW notepad") launch
        # a fresh process — the deterministic new-blank-document path.
        from mini_kio.core.pipeline import _ExecutionCoordinator

        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or
            ({"success": True, "message": "ok"} if cap.startswith("keyboard_") else {"success": False, "message": "no"})
        )
        coord = _ExecutionCoordinator()
        launched = []
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus",
            return_value={"success": True, "message": "Focused Notepad."},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text", return_value=[],
        ), mock.patch.object(
            coord, "_wait_for_blank_new_window", return_value=9002,
        ), mock.patch.object(
            coord, "_read_window_text_for", return_value="",
        ), mock.patch(
            "mini_kio.core.execution_boundary.execute_action",
            side_effect=lambda a, t: launched.append(a) or {"success": True, "message": "opened", "pid": 77},
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "a new notepad", "metadata": {"payload": "hello"}},
                None,
            )
        self.assertTrue(result["success"])
        caps = [c for c, _ in events]
        # a fresh PROCESS is launched (not a ctrl+n hotkey), then typed into.
        self.assertEqual(launched, ["open_app"])
        self.assertIn("keyboard_type", caps)
        self.assertNotIn("keyboard_hotkey", caps)

    def test_type_into_freshly_opened_app_skips_ctrl_n(self):
        # If KIO itself just launched the app, it is already a blank document
        # — no extra Ctrl+N new-document shortcut is needed.
        from mini_kio.core.pipeline import _ExecutionCoordinator

        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or
            ({"success": True, "message": "ok"} if cap.startswith("keyboard_") else {"success": False, "message": "no"})
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider",
            return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus",
            return_value={"success": True, "message": "Focused Notepad."},
        ), mock.patch(
            "mini_kio.core.execution_boundary.execute_action",
            return_value={"success": True, "pid": 4242},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text", return_value=[],
        ), mock.patch.object(
            coord, "_wait_for_blank_new_window", return_value=9003,
        ), mock.patch.object(
            coord, "_read_window_text_for", return_value="",
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": "hello"}},
                None,
            )
        self.assertTrue(result["success"])
        caps = [c for c, _ in events]
        # keyboard_type present; no ctrl+n hotkey issued for a fresh launch.
        self.assertIn("keyboard_type", caps)
        self.assertNotIn("keyboard_hotkey", caps)

    def test_press_enter(self):
        d = _classify("press enter")
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "key_press")
        self.assertEqual(d.target, "enter")

    def test_press_enter(self):
        d = _classify("press enter")
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "key_press")
        self.assertEqual(d.target, "enter")

    def test_press_ctrl_s(self):
        d = _classify("press ctrl+s")
        self.assertEqual(d.action, "key_press")
        self.assertEqual(d.target, "ctrl+s")

    def test_press_release_is_not_key_press(self):
        d = _classify("press release")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_hit_me_with_your_best_shot_is_not_key_press(self):
        d = _classify("hit me with your best shot")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_shortcut_family_is_semantic(self):
        # Capability-quality: edit shortcuts are SEMANTIC actions (SAVE/COPY/
        # PASTE/SELECT_ALL/UNDO/REDO) whose key combo lives in metadata — not
        # mechanical key_press with a hardcoded target. The executor resolves
        # the contextual target and verifies the real result.
        self.assertEqual(_classify("save").action, "save")
        self.assertEqual((_classify("save").metadata or {}).get("combo"), "ctrl+s")
        self.assertEqual(_classify("copy that").action, "copy")
        self.assertEqual((_classify("copy that").metadata or {}).get("combo"), "ctrl+c")
        self.assertEqual(_classify("paste it here").action, "paste")
        self.assertEqual(_classify("select all").action, "select_all")
        self.assertEqual(_classify("undo").action, "undo")
        self.assertEqual(_classify("redo").action, "redo")

    def test_scroll_and_click(self):
        self.assertEqual(_classify("scroll down").action, "scroll")
        self.assertEqual(_classify("scroll to the bottom").target, "bottom")
        d = _classify("click here")
        self.assertEqual(d.action, "click")

    def test_click_named_element_is_not_faked(self):
        d = _classify("click the search box")
        self.assertNotEqual(d.intent_type.value, "desktop_action")

    def test_multi_step_desktop_actions_route_through_canonical_executor(self):
        from mini_kio.core.command_router import _run_desktop_action_step

        with mock.patch(
            "mini_kio.core.pipeline._ExecutionCoordinator._exec_desktop_action",
            return_value={"success": True, "message": "typed"},
        ) as exec_action:
            steps = parse_command("open notepad and type hello and save")
            actions = [s["action"] for s in steps]
            self.assertIn("type", actions)
            self.assertIn("save", actions)
            # type step -> canonical type with payload
            r = _run_desktop_action_step("type", "hello")
            self.assertTrue(r["success"])
            called = exec_action.call_args[0][0]
            self.assertEqual(called["action"], "type")
            self.assertEqual(called["metadata"]["payload"], "hello")
            # save step -> semantic SAVE with the combo in metadata
            r2 = _run_desktop_action_step("save", "")
            self.assertTrue(r2["success"])
            called2 = exec_action.call_args[0][0]
            self.assertEqual(called2["action"], "save")
            self.assertEqual(called2["metadata"]["combo"], "ctrl+s")


class LockUnlockTest(unittest.TestCase):
    """Section 8: lock/unlock is a single authoritative system-session model.
    All variants resolve to the same canonical owner; responses are truthful
    and carry the lock_state."""

    def test_unlock_variants_resolve_to_canonical_owner(self):
        for phrase in ("unlock", "unlock the system", "wake and unlock"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "system", phrase)
            self.assertEqual(d.action, "unlock_system", phrase)

    def test_lock_variants_resolve_to_canonical_owner(self):
        for phrase in ("lock", "lock the computer", "secure my pc"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "system", phrase)
            self.assertEqual(d.action, "lock_system", phrase)

    def test_lock_state_queries_route_operational(self):
        for phrase in ("is my computer locked", "what is the lock status",
                       "can i use the computer", "is the pc unlocked"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "operational", phrase)
            self.assertEqual(d.action, "lock_state", phrase)

    def test_lock_state_returns_truthful_dict(self):
        from mini_kio.core.system_operator import lock_state
        result = lock_state()
        self.assertIn("lock_state", result)
        self.assertIn(result["lock_state"], ("locked", "unlocked", "unknown"))
        if result["lock_state"] == "locked":
            self.assertIn("locked", result["message"].lower())
        elif result["lock_state"] == "unlocked":
            self.assertIn("isn't locked", result["message"].lower())

    def test_unlock_never_claims_unlocked_without_state(self):
        # A LOCKED workstation must yield authentication_required, never a
        # fabricated UNLOCKED.
        from mini_kio.core import system_operator
        with mock.patch.object(system_operator, "is_workstation_locked", return_value=True):
            result = system_operator.unlock_system()
            self.assertEqual(result["lock_state"], "authentication_required")
            self.assertNotIn("isn't locked", result["message"].lower())
        with mock.patch.object(system_operator, "is_workstation_locked", return_value=False):
            result = system_operator.unlock_system()
            self.assertEqual(result["lock_state"], "unlocked")


class SystemNounBoundaryTest(unittest.TestCase):
    """System nouns stay OUT of the close-verb family."""

    def test_shut_down_the_computer_is_shutdown(self):
        d = _classify("shut down the computer")
        self.assertEqual(d.intent_type.value, "system")
        self.assertEqual(d.action, "shutdown_system")

    def test_shut_down_my_pc_is_shutdown(self):
        d = _classify("shut down my pc")
        self.assertEqual(d.action, "shutdown_system")


class DegradedScopedGatingTest(unittest.TestCase):
    """Live-found (2026-08-12): a browser connector outage escalated the
    WHOLE runtime to DEGRADED and blocked unrelated native app launch/close.
    Gate 2.5 refinement: integrity degrades per capability group, and the
    safety gate blocks only the degraded group's actions."""

    def test_browser_outage_degrades_browser_group_only(self):
        from mini_kio.core import runtime as rt
        from mini_kio.core.execution_boundary import check_safety_policy

        r = rt.KioRuntime()
        rt._CURRENT_RUNTIME = r
        # 3 browser_goto failures (operator_reported_failure -> medium, w=2)
        # cross the degraded threshold (6) for the browser group.
        for _ in range(3):
            rt.record_runtime_integrity_warning(
                "execution_failure",
                {"action": "browser_goto", "target": "https://x.com",
                 "failure_class": "operator_reported_failure"},
            )
        self.assertEqual(r.safety_state, "DEGRADED")
        self.assertEqual(r.degraded_capability_groups, {"browser"})

        # Browser-group actions are blocked...
        allowed, reason = check_safety_policy("browser_goto", "https://x.com", r)
        self.assertFalse(allowed)
        self.assertIn("browser", reason)
        allowed, _ = check_safety_policy(
            "execute_capability", "chrome::open_url::https://chatgpt.com::chatgpt", r)
        self.assertFalse(allowed)

        # ...but native launch/close, media, and system control stay available.
        allowed, _ = check_safety_policy("open_app", "notepad", r)
        self.assertTrue(allowed)
        allowed, _ = check_safety_policy("close_app", "notepad", r)
        self.assertTrue(allowed)
        allowed, _ = check_safety_policy("open_folder", "downloads", r)
        self.assertTrue(allowed)
        allowed, _ = check_safety_policy("lock_system", "", r)
        self.assertTrue(allowed)
        allowed, _ = check_safety_policy("media_play", "", r)
        self.assertTrue(allowed)

    def test_benign_failures_never_degrade_runtime(self):
        # "open <nonexistent>", "close <not running>", invalid URLs, etc.
        # are normal user-request outcomes, NOT runtime integrity problems.
        from mini_kio.core import runtime as rt

        r = rt.KioRuntime()
        rt._CURRENT_RUNTIME = r
        for _ in range(12):
            rt.record_runtime_integrity_warning(
                "execution_failure",
                {"action": "open_app", "target": "nonexistentapp",
                 "failure_class": "not_installed"},
            )
        self.assertEqual(r.safety_state, "NORMAL")
        self.assertEqual(r.integrity_score, 0)
        self.assertEqual(r.degraded_capability_groups, set())

    def test_group_scores_reset_on_manual_recovery(self):
        from mini_kio.core import runtime as rt

        r = rt.KioRuntime()
        rt._CURRENT_RUNTIME = r
        for _ in range(3):
            rt.record_runtime_integrity_warning(
                "execution_failure",
                {"action": "browser_goto", "target": "https://x.com",
                 "failure_class": "operator_reported_failure"},
            )
        self.assertEqual(r.safety_state, "DEGRADED")
        rt.manual_runtime_recovery()
        self.assertEqual(r.safety_state, "NORMAL")
        self.assertEqual(r.integrity_group_scores, {})
        self.assertEqual(r.degraded_capability_groups, set())


class TargetInstanceSemanticsTest(unittest.TestCase):
    """Section 2/3: the semantic layer must distinguish open / open another /
    new tab / new window / new instance / focus — generically, never per-app.
    Entity identity is preserved while INSTANCE changes."""

    def test_plain_open_is_existing_target(self):
        d = _classify("open chatgpt")
        self.assertFalse(d.metadata.get("explicit_new"))

    def test_open_another_tab_marks_explicit_new(self):
        d = _classify("open another tab of chatgpt")
        self.assertEqual(d.intent_type.value, "browser_navigate")
        self.assertEqual(d.action, "execute_capability")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.metadata.get("instance"), "tab")
        self.assertIn("chatgpt", d.target)

    def test_open_a_new_entity_tab_trailing_form(self):
        # Live-found semantic gap (2026-08-12): "open a new chatgpt tab" was
        # routed to NATIVE open_app with target "chatgpt tab" (the trailing
        # "tab" leaked into the entity). It must be a NEW browser-tab
        # instance of the entity, never a native app named "chatgpt tab".
        d = _classify("open a new chatgpt tab")
        self.assertEqual(d.intent_type.value, "browser_navigate")
        self.assertEqual(d.action, "execute_capability")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.metadata.get("instance"), "tab")
        self.assertIn("chatgpt", d.target)
        self.assertNotIn("tab", d.target.split("::")[-1])  # entity is chatgpt, not chatgpt tab

    def test_open_another_entity_tab_trailing_form(self):
        d = _classify("open another github tab")
        self.assertEqual(d.action, "execute_capability")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.metadata.get("instance"), "tab")
        self.assertIn("github", d.target)

    def test_open_in_new_tab_has_browser_prefix(self):
        # Live-found latent bug: "open X in a new tab" produced an EMPTY
        # browser prefix ("::open_url::...") which broke execute_capability
        # parsing. The modal target must carry the resolved default browser.
        d = _classify("open chatgpt in a new tab")
        self.assertTrue(d.metadata.get("explicit_new"))
        prefix = d.target.split("::")[0]
        self.assertIn(prefix, ("chrome", "edge", "firefox", "brave", "comet"))

    def test_new_browser_window_is_window_instance(self):
        d = _classify("open a new browser window for github")
        self.assertEqual(d.intent_type.value, "browser_navigate")
        self.assertEqual(d.action, "execute_capability")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.metadata.get("instance"), "window")
        self.assertIn("github", d.target)

    def test_another_entity_window_is_window_instance(self):
        d = _classify("open another github window")
        self.assertEqual(d.action, "execute_capability")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.metadata.get("instance"), "window")

    def test_open_in_new_window_is_window_instance(self):
        d = _classify("open github in a new window")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.metadata.get("instance"), "window")

    def test_open_in_new_tab_marks_explicit_new(self):
        d = _classify("open chatgpt in a new tab")
        self.assertTrue(d.metadata.get("explicit_new"))

    def test_open_another_instance_marks_explicit_new(self):
        d = _classify("open another calculator")
        self.assertEqual(d.intent_type.value, "desktop_open")
        self.assertEqual(d.action, "open_app")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.target, "calculator")

    def test_open_a_new_instance_marks_explicit_new(self):
        d = _classify("open a new calculator")
        self.assertTrue(d.metadata.get("explicit_new"))
        self.assertEqual(d.target, "calculator")

    def test_focus_is_not_explicit_new(self):
        d = _classify("focus chatgpt")
        self.assertEqual(d.intent_type.value, "browser_focus")
        self.assertEqual(d.action, "focus")

    def test_switch_to_is_focus_not_new(self):
        d = _classify("switch to chatgpt")
        self.assertEqual(d.action, "focus")

    def test_plain_open_skips_force_new_dup_prevention(self):
        # The executor path: default open must go through duplicate
        # prevention; explicit-new opens must skip it (force_new).
        from mini_kio.core import app_operator
        from mini_kio.core import command_router

        with mock.patch.object(
            app_operator, "_find_existing_web_target", return_value=None
        ), mock.patch.object(
            command_router, "_get_connector", return_value=None
        ):
            # friendly_name with ::new marker -> force_new skips dedup
            result = app_operator.execute_capability(
                "chrome::open_url::https://chatgpt.com::chatgpt::new")
        self.assertIn("chatgpt", result.get("message", "").lower())


class LLMBypassAuditTest(unittest.TestCase):
    """Section 7 (2026-08-12 directive): deterministic local/system state
    queries must NEVER fall through to LLM/web knowledge. The app inventory
    is answered from actual Windows state, whatever the possession wording."""

    def test_what_apps_do_you_have_is_inventory(self):
        # Live-found: "what apps do you have" produced irrelevant external
        # web content. It is the SAME inventory as "what apps do I have".
        for phrase in (
            "what apps do you have",
            "what apps does the system have",
            "what apps does the computer have",
            "what apps do i have",
            "what apps are installed",
            "what apps have you got",
            "what software is installed",
            "what programs are installed",
            "what apps are on my computer",
        ):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "operational", phrase)
            self.assertEqual(d.action, "app_inventory", phrase)

    def test_installed_existence_query_is_os_probe(self):
        for phrase in ("is notepad installed", "is winrar installed"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "operational", phrase)
            self.assertEqual(d.action, "app_installed", phrase)

    def test_informal_kio_possessive_is_operational(self):
        # Live-found leak: "whats kios status" (informal possessive) reached
        # the identity path instead of the deterministic operational status
        # owner. "kios" must canonicalize to "kio's" for the KIO-self
        # families, while genuinely different words ("kiosk") stay on the
        # knowledge path.
        for phrase in ("whats kios status", "kios health", "what is kios uptime", "kios status"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "operational", phrase)
        self.assertNotIn(
            d := _classify("what is a kiosk").intent_type.value, ("operational", "identity")
        )

    def test_running_state_query_is_os_probe(self):
        d = _classify("is spotify running")
        self.assertEqual(d.intent_type.value, "operational")
        self.assertEqual(d.action, "app_running")

    def test_whats_running_on_computer_is_desktop_state(self):
        # Live-found (2026-08-12 revalidation): "what's running on my
        # computer" fell through to web knowledge and described the Android
        # app "WhatsRunning". The whole running-state family must route to the
        # deterministic desktop-state owner (list_tabs), whatever the trailing
        # location or the noun (apps/processes).
        for phrase in (
            "what's running on my computer",
            "what's running on this pc",
            "what is running on your system",
            "what's running on the machine",
            "what processes are running",
            "which processes are running",
            "what's running right now",
            "what apps are running",
            "what's currently running",
        ):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "browser_tabs", phrase)
            self.assertEqual(d.action, "list_tabs", phrase)
        # Knowledge shapes that merely share a word must NOT be hijacked.
        for phrase in ("what is running time", "what's open source",
                       "how much does ram cost"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "information", phrase)

    def test_system_healthy_wording_routes_system(self):
        # "is the system healthy" must agree with "is my computer okay":
        # deterministic system owner, never the conversation fallback.
        for phrase in ("is the system healthy", "is my computer okay",
                       "is the computer healthy"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "operational", phrase)
            self.assertEqual(d.action, "system", phrase)


class NewWindowTruthfulnessTest(unittest.TestCase):
    """Live-found (2026-08-12 revalidation): "open github in a new window"
    replied "Opened Github: in a new Chrome window." — (1) the "::newwindow"
    marker strip used [:-10] on an 11-char marker, leaving a trailing colon in
    the friendly name, and (2) the new-window claim carried verification_mode
    "noop" (no evidence). The canonical owner must strip the full marker and
    only claim a separate window when one is actually observed."""

    def test_newwindow_strip_and_window_verification(self):
        from types import SimpleNamespace
        from mini_kio.core import app_operator as ao
        from mini_kio.core import routing_utils as ru
        with mock.patch.object(ao, "_find_in_registry",
                               return_value={"exe": "chrome.exe",
                                             "process": "chrome.exe"}), \
                mock.patch.object(ao, "_resolve_path",
                                  return_value=r"C:\nonexistent\chrome.exe"), \
                mock.patch.object(ao, "subprocess") as sp, \
                mock.patch.object(ao, "_visible_browser_window_count",
                                  side_effect=[0, 1]), \
                mock.patch.object(ru, "register_browser_capability"):
            sp.Popen.return_value = SimpleNamespace(pid=12345)
            result = ao.execute_capability(
                "chrome::open_url::https://github.com::github::newwindow"
            )
        # Full marker stripped: no "Github:" colon artifact, clean display.
        self.assertNotIn("Github:", result["message"])
        self.assertIn("GitHub", result["message"])
        # New window actually observed (window count 0 -> 1): verified claim.
        self.assertEqual(result.get("verification_mode"), "window_identity")

    def test_newwindow_unconfirmed_is_truthful(self):
        # If no extra window is observed, KIO must NOT claim a new window.
        from types import SimpleNamespace
        from mini_kio.core import app_operator as ao
        from mini_kio.core import routing_utils as ru
        with mock.patch.object(ao, "_find_in_registry",
                               return_value={"exe": "chrome.exe",
                                             "process": "chrome.exe"}), \
                mock.patch.object(ao, "_resolve_path",
                                  return_value=r"C:\nonexistent\chrome.exe"), \
                mock.patch.object(ao, "subprocess") as sp, \
                mock.patch.object(ao, "_visible_browser_window_count",
                                  side_effect=[1, 1, 1, 1, 1]), \
                mock.patch.object(ru, "register_browser_capability"):
            sp.Popen.return_value = SimpleNamespace(pid=12345)
            result = ao.execute_capability(
                "chrome::open_url::https://github.com::github::newwindow"
            )
        self.assertEqual(result.get("verification_mode"), "window_unverified")
        self.assertNotIn("new Chrome window", result["message"])
        self.assertIn("couldn't confirm", result["message"])

    def test_knowledge_queries_stay_knowledge(self):
        # "what is notepad" must stay a KNOWLEDGE query, never an OS action.
        d = _classify("what is notepad")
        self.assertEqual(d.intent_type.value, "information")


class CompanionIntelligenceAuditTest(unittest.TestCase):
    """Section 17 (2026-08-12 directive): Companion Intelligence — the
    doctrine's dimensions (independent judgment, preference, disagreement,
    recommendation, uncertainty, curiosity, emotional modelling, self-
    evaluation) must route to the conversational/identity owners, NEVER to
    web knowledge or the wrong canonical owner."""

    def test_opinion_judgment_routes_conversation(self):
        # Curated opinion topics ("ai", "technology") have deterministic
        # canonical answers; uncurated topics route to the conversational LLM.
        for phrase in (
            "what is your take on remote work",
            "in your opinion, is it worth it",
            "what do you think about urban farming",
        ):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "conversation", phrase)
            self.assertEqual(d.action, "converse", phrase)
        # Curated deterministic opinions stay on the identity owner.
        for phrase in ("what do you think about ai", "what are your opinions on technology"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "identity", phrase)


    def test_preference_and_recommendation_route_conversation(self):
        for phrase in (
            "do you like jazz music",
            "what is your favorite color",
            "should i watch dune",
            "recommend a good movie for tonight",
            "would you recommend that book",
            "which is better, apple or android",
        ):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "conversation", phrase)
            self.assertEqual(d.action, "converse", phrase)

    def test_converse_prompt_allows_reasoned_preferences(self):
        # Live-found (2026-08-12): "do you like jazz music" deflected with
        # "As an AI, I don't have personal preferences" — doctrine's Section 5
        # forbids using "I'm an AI" as a substitute for truthful behavior. The
        # converse system prompt must permit reasoned preferences/curiosity
        # while still forbidding fabricated human experience.
        from mini_kio.llm.llm_constants import _ASYSTEM_PROMPT
        joined = _ASYSTEM_PROMPT.lower()
        self.assertIn("i'd choose", joined)
        self.assertIn("deflecting", joined)
        self.assertIn("never fabricate human memories", joined)
        self.assertIn("curiosity", joined)

    def test_disagreement_routes_conversation(self):
        d = _classify("i disagree with you")
        self.assertEqual(d.intent_type.value, "conversation")

    def test_emotional_expression_routes_empathy(self):
        for phrase in ("i feel really stressed today", "im so excited"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "conversation", phrase)
            self.assertEqual(d.action, "empathy", phrase)

    def test_curiosity_is_conversation_not_static_identity(self):
        # Live-found: "what are you curious about" matched the identity
        # dataset's broad "what are you" prefix and returned the static
        # "who are you" identity. Curiosity is a conversational dimension
        # (doctrine Section 5) and must route to the LLM with personality
        # context, while "who are you" stays the deterministic identity.
        for phrase in (
            "what are you curious about",
            "what interests you",
            "what excites you",
            "what are you interested in",
        ):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "conversation", phrase)
            self.assertEqual(d.action, "converse", phrase)
        # Identity questions remain deterministic.
        for phrase in ("who are you", "what is kio", "what are you"):
            d = _classify(phrase)
            self.assertEqual(d.intent_type.value, "identity", phrase)


class MCPStartupNonBlockingTest(unittest.TestCase):
    """Section 9: optional MCP provider initialization must not block the
    deterministic local command path at startup."""

    def test_start_returns_immediately(self):
        from mini_kio.core.mcp.registry import MCPServerRegistry, MCPServerInfo
        from mini_kio.core.mcp.client import MCPClient

        registry = MCPServerRegistry()
        for stype in ("filesystem", "git", "terminal", "sqlite", "docker",
                      "github", "postgres", "redis"):
            registry.register_server(MCPServerInfo(server_type=stype, display_name=stype))

        blocked = mock.Mock(side_effect=lambda info: time.sleep(0.5))
        client = MCPClient(registry)
        with mock.patch.object(client, "_connect", blocked):
            start = time.monotonic()
            client.start()
            elapsed = time.monotonic() - start
        # start() must return without waiting for any server handshake.
        self.assertLess(elapsed, 0.3)
        self.assertEqual(blocked.call_count, 8)
        time.sleep(0.9)  # allow the daemon threads to finish


class BrowserModalityGatingTest(unittest.TestCase):
    """Live-found (2026-08-12 revalidation): "open youtube in comet" opened a
    Chrome tab AND launched Comet — the Chrome extension connector served an
    explicit non-default-browser request. The connector may only serve the
    DEFAULT browser; explicit edge/comet/firefox/brave requests must go
    straight to their own browser binary (never a silent Chrome substitution).
    """

    def _make_conn(self):
        from types import SimpleNamespace
        conn = mock.Mock()
        conn.is_connected.return_value = True
        conn.open_tab = mock.AsyncMock(return_value=SimpleNamespace(
            success=True, tab_id="t1", url="https://youtube.com",
            title="YouTube", window_id=1,
        ))
        return conn

    def _patch_env(self, default_browser):
        from mini_kio.core import app_operator as ao
        from mini_kio.core import routing_utils as ru
        from mini_kio.core import command_router as cr
        from mini_kio.core import config as cfg
        from unittest import mock
        return [
            mock.patch.object(ao, "_find_in_registry",
                              return_value={"exe": "comet.exe",
                                            "process": "comet.exe"}),
            mock.patch.object(ao, "_resolve_path",
                              return_value=r"C:\fake\comet.exe"),
            mock.patch.object(ao, "subprocess"),
            mock.patch.object(ao, "_refine_pid_windows", return_value=9876),
            mock.patch.object(ru, "register_browser_capability"),
            mock.patch.object(cr, "_get_connector"),
            mock.patch.object(cfg, "DEFAULT_BROWSER", default_browser),
        ]

    def test_nondefault_browser_skips_connector(self):
        # DEFAULT_BROWSER=chrome, request is comet: the Chrome connector must
        # NOT be used (no Chrome tab side effect); Comet binary is launched.
        from types import SimpleNamespace
        from mini_kio.core import app_operator as ao
        conn = self._make_conn()
        patches = self._patch_env("chrome")
        with patches[6], patches[5] as get_conn, patches[0], patches[1], \
                patches[2] as sp, patches[3], patches[4]:
            get_conn.return_value = conn
            sp.Popen.return_value = SimpleNamespace(pid=999)
            result = ao.execute_capability(
                "comet::open_url::https://youtube.com::youtube"
            )
        conn.open_tab.assert_not_called()
        self.assertIn("Comet", result["message"])
        self.assertTrue(result["success"])

    def test_default_browser_uses_connector(self):
        # DEFAULT_BROWSER=chrome, request is chrome: the connector path runs
        # (tab-level open with verification), no binary relaunch.
        from mini_kio.core import app_operator as ao
        conn = self._make_conn()
        patches = self._patch_env("chrome")
        with patches[6], patches[5] as get_conn, patches[0], patches[1], \
                patches[2] as sp, patches[3], patches[4]:
            get_conn.return_value = conn
            result = ao.execute_capability(
                "chrome::open_url::https://youtube.com::youtube"
            )
        conn.open_tab.assert_awaited_once()
        self.assertIn("YouTube", result["message"])



class InstanceMarkerReferentTest(unittest.TestCase):
    """Live-found (post-commit 84a95b0 smoke): after "open a new chatgpt tab"
    the context referent became 'new' (the ::new INSTANCE marker was captured
    as the entity by parse_target's lazy arg + trailing name group), so a later
    "close it" resolved to 'close new' -> truthful but useless failure.
    The entity (chatgpt) must stay the referent; the instance marker is a
    qualifier, never a name.
    """

    def test_newwindow_never_becomes_referent(self):
        from mini_kio.core.target_ref import parse_target, safe_target_name
        ref = parse_target("chrome::open_url::https://github.com::github::newwindow")
        self.assertEqual(ref.name, "github")
        self.assertEqual(ref.kind, "webapp")
        self.assertEqual(ref.url, "https://github.com")
        self.assertEqual(safe_target_name(
            "chrome::open_url::https://github.com::github::newwindow"), "github")

    def test_newtab_never_becomes_referent(self):
        from mini_kio.core.target_ref import parse_target, safe_target_name
        ref = parse_target("chrome::open_url::https://chat.openai.com::chatgpt::new")
        self.assertEqual(ref.name, "chatgpt")
        self.assertEqual(ref.url, "https://chat.openai.com")
        self.assertEqual(safe_target_name(
            "chrome::open_url::https://chat.openai.com::chatgpt::new"), "chatgpt")

    def test_plain_capability_name_unchanged(self):
        # No marker: name behavior is untouched.
        from mini_kio.core.target_ref import parse_target, safe_target_name
        ref = parse_target("chrome::open_url::https://chat.openai.com::chatgpt")
        self.assertEqual(ref.name, "chatgpt")
        self.assertEqual(safe_target_name(
            "chrome::open_url::https://chat.openai.com::chatgpt"), "chatgpt")

    def test_context_close_it_resolves_to_entity_not_marker(self):
        # End-to-end referent: after a NEW-TAB execute_capability the
        # context manager must remember 'chatgpt', so 'close it' becomes
        # 'close chatgpt' — never 'close new'.
        from mini_kio.core.context_manager import SessionContext
        from mini_kio.core.target_ref import safe_target_name
        cm = SessionContext(session_id="test-session")
        target = "chrome::open_url::https://chat.openai.com::chatgpt::new"
        result = {
            "success": True,
            "action": "execute_capability",
            "target": target,
        }
        cm.update(result, "open a new chatgpt tab")
        self.assertEqual(cm.active_entity, safe_target_name(target))
        self.assertEqual(cm.active_entity, "chatgpt")
        resolved = cm.resolved_text("close it")
        self.assertEqual(resolved, "close chatgpt")


class CompanionPreferenceTest(unittest.TestCase):
    """Doctrine-aligned companion behavior: KIO has stable MODELED preferences
    and may express concrete reasoned choices — never the deflecting
    "I'm an AI, so I don't have preferences", and never fabricated human
    biography (childhood, senses, memories, lived experience).
    """

    def test_modeled_preferences_are_stable_character_data(self):
        from mini_kio.llm.KIO_character_knowledge import resolve_modeled_preferences
        prefs = resolve_modeled_preferences()
        self.assertTrue(len(prefs) >= 5)
        # Stable across calls — same object identity semantics, no random flip.
        self.assertEqual(resolve_modeled_preferences(), prefs)
        # Character values, not biography claims.
        joined = " ".join(prefs).lower()
        self.assertNotIn("childhood", joined)
        self.assertNotIn("i remember", joined)
        self.assertIn("prefer", joined)

    def test_converse_prompt_allows_reasoned_preferences(self):
        # The live converse prompt must ALLOW concrete reasoned choices and
        # forbid the "I'm an AI, no preferences" deflection — the previous
        # text forced analytic deflection which produced that exact failure.
        import mini_kio.core.pipeline as pl
        src = open(pl.__file__, encoding="utf-8").read()
        # The ban on deflection is present...
        self.assertIn("deflecting with", src)
        # ...and the old force-analytic-only clause is gone.
        self.assertNotIn("discuss them analytically instead", src)
        # Modeled preferences are injected from the canonical character layer.
        self.assertIn("resolve_modeled_preferences", src)
        self.assertIn("stable modeled preferences", src)
        # The anti-fabrication boundary must survive alongside the allowance
        # (assert on single source lines — the phrase spans a string split).
        self.assertIn("never invent ", src)
        self.assertIn("a childhood, ", src)
        self.assertIn("memories, or lived experience", src)

    def test_preference_questions_route_to_converse(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        for q in (
            "do you like jazz music",
            "what's your favorite movie",
            "would you rather use python or javascript",
            "which approach would you choose",
            "what would you recommend",
            "do you agree with that",
            "what are you curious about",
        ):
            d = p._classifier.classify(q, q)
            self.assertEqual(
                d.intent_type.value, "conversation", f"{q!r} should be conversation"
            )
            self.assertEqual(d.action, "converse", f"{q!r} action")

    def test_anti_fabrication_boundary_preserved(self):
        # The doctrine boundary: modeled preference yes, invented biography no.
        from mini_kio.llm.KIO_character_knowledge import (
            resolve_canonical_truth,
            resolve_anti_hallucination_rules,
        )
        t = resolve_canonical_truth("not_human")
        self.assertIsNotNone(t)
        rules = " ".join(resolve_anti_hallucination_rules()).lower()
        self.assertIn("memory", rules)



class GeneratedTypeTest(unittest.TestCase):
    """Capability-quality: TYPE must distinguish literal text from generated
    content. "type hello into notepad" types verbatim; "write a short poem
    about X in notepad" requests CONTENT GENERATION (LLM) then deterministic
    execution. A bare referent pronoun with no context asks for clarification.
    """

    def test_literal_payload_is_not_generated(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify("type hello into notepad", "type hello into notepad")
        self.assertEqual(d.action, "type")
        self.assertFalse((d.metadata or {}).get("generate"))
        self.assertEqual((d.metadata or {}).get("payload"), "hello")

    def test_poem_request_is_generated_content(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "write a short poem comparing messi and ronaldo in notepad",
            "write a short poem comparing messi and ronaldo in notepad",
        )
        self.assertEqual(d.action, "type")
        self.assertTrue((d.metadata or {}).get("generate"))
        self.assertEqual(d.target, "notepad")

    def test_n_point_summary_is_generated(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "put a three-point study summary into notepad",
            "put a three-point study summary into notepad",
        )
        self.assertEqual(d.action, "type")
        self.assertTrue((d.metadata or {}).get("generate"))

    def test_bare_pronoun_requests_clarification(self):
        # A "this" that survived referent resolution must NOT be typed literally.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify("write this in notepad", "write this in notepad")
        self.assertEqual(d.action, "type")
        self.assertEqual((d.metadata or {}).get("payload"), "this")
        # Executor guard: ask for content instead of typing the word.
        from mini_kio.core import pipeline as pl
        from unittest import mock
        with mock.patch.object(pl._ExecutionCoordinator, "_try_native_focus",
                               return_value={"success": True}):
            exec_result = pl._ExecutionCoordinator()._exec_desktop_action(
                {"action": "type", "target": "notepad",
                 "metadata": {"payload": "this"}},
                None,
            )
        self.assertFalse(exec_result["success"])
        self.assertIn("What should I type", exec_result["message"])


class SemanticEditActionTest(unittest.TestCase):
    """Capability-quality: SAVE / COPY / PASTE / SELECT_ALL / UNDO / REDO are
    semantic operations with contextual target resolution + real verification
    — not mechanical keypresses with unverified success messages."""

    def _make_coord(self, events, dp):
        from mini_kio.core.pipeline import _ExecutionCoordinator
        return _ExecutionCoordinator()

    def _typed_events(self):
        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or
            ({"success": True, "message": "ok", "text": "hello world"} if cap == "clipboard_get" else
             {"success": True, "message": "ok"} if cap.startswith("keyboard_") else
             {"success": False, "message": "no"})
        )
        return events, dp

    def test_copy_is_verified_against_real_clipboard(self):
        # COPY reads the clipboard back and reports the ACTUAL copied content
        # length — a truthful verification, not a bare "Pressed Ctrl+C."
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events, dp = self._typed_events()
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(coord, "_try_native_focus", return_value=None), \
             mock.patch.object(coord, "_resolve_edit_target", return_value=None):
            result = coord._exec_desktop_action(
                {"action": "copy", "target": "", "metadata": {"combo": "ctrl+c"}},
                None,
            )
        self.assertTrue(result["success"])
        self.assertIn("11 characters", result["message"])
        self.assertTrue(result.get("verified"))
        self.assertIn(("keyboard_hotkey", "ctrl+c"), events)
        self.assertIn(("clipboard_get", ""), events)

    def test_paste_with_empty_clipboard_is_truthful(self):
        # PASTE with an empty clipboard is a truthful failure, never a fake
        # "Pasted."
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events, dp = self._typed_events()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or
            {"success": False, "message": "clipboard empty"} if cap == "clipboard_get" else
            {"success": True, "message": "ok"}
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(coord, "_try_native_focus", return_value=None), \
             mock.patch.object(coord, "_resolve_edit_target", return_value=None):
            result = coord._exec_desktop_action(
                {"action": "paste", "target": "", "metadata": {"combo": "ctrl+v"}},
                None,
            )
        self.assertFalse(result["success"])
        self.assertIn("empty", result["message"].lower())

    def test_save_resolves_contextual_target_before_key(self):
        # SAVE resolves the session's active entity and focuses it BEFORE
        # sending Ctrl+S — the key lands in the window the user is working in.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events, dp = self._typed_events()
        coord = _ExecutionCoordinator()
        focused = []
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(
            coord, "_resolve_edit_target", return_value="notepad",
        ), mock.patch.object(
            coord, "_try_native_focus",
            side_effect=lambda t: focused.append(t) or {"success": True, "message": "focused"},
        ):
            result = coord._exec_desktop_action(
                {"action": "save", "target": "", "metadata": {"combo": "ctrl+s"}},
                None,
            )
        self.assertTrue(result["success"])
        self.assertEqual(focused, ["notepad"])
        self.assertIn(("keyboard_hotkey", "ctrl+s"), events)
        self.assertIn("Saved", result["message"])

    def test_select_all_is_semantic_action(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify("select everything", "select everything")
        self.assertEqual(d.action, "select_all")
        self.assertEqual((d.metadata or {}).get("combo"), "ctrl+a")

    def test_type_verification_confirmed_payload(self):
        # TYPE read-back confirms the payload actually landed in the window.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or {"success": True, "message": "ok"}
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus",
            return_value={"success": True, "message": "Focused Notepad."},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text",
            return_value=["hello there my friend"],
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": "hello"}},
                None,
            )
        self.assertTrue(result["success"])
        self.assertTrue(result.get("verified"))
        # Natural outcome wording — never internal "verified" vocabulary.
        self.assertNotIn("verified", result["message"].lower())
        self.assertTrue("done" in result["message"].lower() or "typed" in result["message"].lower())

    def test_large_payload_uses_atomic_paste_not_per_char_typing(self):
        # Root-cause regression: generated documents were typed one character
        # at a time (slow, interruptible, truncation-prone → mixed/partial
        # content). Large payloads MUST use atomic clipboard paste, with
        # keyboard_type only as fallback for short literals.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events = []
        dp = mock.MagicMock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or {"success": True, "message": "ok"}
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus", return_value={"success": True},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text", return_value=[]
        ), mock.patch.object(
            coord, "_snapshot_app_windows", return_value={10, 20},
        ), mock.patch.object(
            coord, "_wait_for_new_window", return_value=30,
        ):
            big = "word " * 60  # ~300 chars > 80 threshold
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": big}},
                None,
            )
        caps = [c for c, _ in events]
        self.assertIn("clipboard_set", caps)
        self.assertIn("keyboard_hotkey", caps)
        self.assertNotIn("keyboard_type", caps)
        self.assertTrue(result["success"])

    def test_short_literal_uses_direct_typing(self):
        # "type hello KIO" is a LITERAL — no content generation, no clipboard
        # round-trip; direct typing is the correct cheap path.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events = []
        dp = mock.MagicMock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or {"success": True, "message": "ok"}
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus", return_value={"success": True},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text", return_value=[]
        ), mock.patch.object(
            coord, "_snapshot_app_windows", return_value=set(),
        ), mock.patch.object(
            coord, "_wait_for_new_window", return_value=5,
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": "hello KIO"}},
                None,
            )
        caps = [c for c, _ in events]
        self.assertIn("keyboard_type", caps)
        self.assertNotIn("clipboard_set", caps)
        self.assertTrue(result["success"])

    def test_type_verification_absent_payload_is_truthful_failure(self):
        # Window WAS readable and the payload is NOT present → truthful
        # failure, never a success-shaped "Typed it."
        from mini_kio.core.pipeline import _ExecutionCoordinator
        events = []
        dp = mock.Mock()
        dp.execute.side_effect = lambda cap, target="": (
            events.append((cap, target)) or {"success": True, "message": "ok"}
        )
        coord = _ExecutionCoordinator()
        with mock.patch(
            "mini_kio.desktop.DesktopProvider", return_value=dp,
        ), mock.patch.object(
            coord, "_try_native_focus",
            return_value={"success": True, "message": "Focused Notepad."},
        ), mock.patch.object(
            coord, "_read_target_or_foreground_text",
            return_value=["completely different content"],
        ):
            result = coord._exec_desktop_action(
                {"action": "type", "target": "notepad", "metadata": {"payload": "hello"}},
                None,
            )
        self.assertFalse(result["success"])
        self.assertIn("isn't showing", result["message"])


class LiveFoundTypeRoutingTest(unittest.TestCase):
    """Live Telegram regressions from the real conversation: content-request
    phrasings that previously fell to companion chat (fabricating 'I'll type
    CTRL+C into Notepad') or parsed the target wrong ('a new notepad file')
    must converge on the canonical TYPE/create_document semantics."""

    def test_do_a_comparison_and_type_it_into_notepad(self):
        # "Do a small comparison on X and Y and type it into notepad" must be
        # a GENERATED-content TYPE task — not a companion preference answer.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "do a small comparison on KTU 2024 scheme and KTU 2019 scheme and type it into notepad",
            "do a small comparison on KTU 2024 scheme and KTU 2019 scheme and type it into notepad",
        )
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "type")
        self.assertEqual(d.target, "notepad")
        self.assertTrue((d.metadata or {}).get("generate"))
        payload = (d.metadata or {}).get("payload", "")
        self.assertIn("comparison", payload)
        self.assertIn("ktu", payload)

    def test_type_the_comparison_onto_notepad(self):
        # "onto" connector + artifact-only referent must classify as TYPE.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "type the comparison onto notepad",
            "type the comparison onto notepad",
        )
        self.assertEqual(d.action, "type")
        self.assertEqual(d.target, "notepad")

    def test_comparison_in_new_notepad_file_normalizes_target(self):
        # "a new notepad file" -> target "notepad" (artifact suffix stripped),
        # never a broken "A new notepad file" target.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "type a comparison on messi vs lewis hamilton in a new notepad file",
            "type a comparison on messi vs lewis hamilton in a new notepad file",
        )
        self.assertEqual(d.action, "type")
        self.assertEqual(d.target, "notepad")
        self.assertTrue((d.metadata or {}).get("generate"))
        self.assertTrue((d.metadata or {}).get("new_instance"))

    def test_doubled_connector_into_onto(self):
        # "do a real comparison and type into onto a new notepad file" — the
        # doubled connector must not break classification.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "do a real comparison and type into onto a new notepad file",
            "do a real comparison and type into onto a new notepad file",
        )
        self.assertEqual(d.action, "type")
        self.assertEqual(d.target, "notepad")
        self.assertTrue((d.metadata or {}).get("generate"))

    def test_bare_write_essay_routes_to_document_creation(self):
        # User directive: "write an essay on X" (no explicit app) routes to
        # LLM content + real artifact creation, not a companion answer.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify("write an essay on KIO", "write an essay on KIO")
        self.assertEqual(d.action, "create_document")
        self.assertEqual(d.target, "kio")

    def test_artifact_only_payload_inherits_topic_from_context(self):
        # "do a real comparison" after "a comparison on messi vs lewis
        # hamilton" inherits the topic from recent context instead of asking.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        from mini_kio.core.pipeline import RoutingDecision
        from mini_kio.core.pipeline import IntentType
        from mini_kio.core.context_manager import get_session_context

        ctx = get_session_context("inherit-topic-test")
        ctx.append_exchange(
            "type a comparison on messi vs lewis hamilton in a new notepad file",
            "Typed it into a new Notepad.",
        )
        coord = _ExecutionCoordinator()
        decision = RoutingDecision(
            IntentType.DESKTOP_ACTION, "type", "notepad",
            "do a real comparison and type into onto a new notepad file",
            "do a real comparison and type into onto a new notepad file",
            confidence=1.0, session_id="inherit-topic-test",
        )
        topic = coord._inherit_content_topic(decision, "comparison")
        self.assertIsNotNone(topic)
        self.assertIn("messi", topic)


class DocumentCreationTest(unittest.TestCase):
    """Capability-quality: CREATE DOCUMENT produces a meaningful artifact with
    a real filename, clean structured content, and verification facts. The
    document operator is the canonical owner (dependency-free OOXML)."""

    def test_meaningful_filenames(self):
        from mini_kio.core.document_operator import generate_document_filename
        self.assertEqual(
            generate_document_filename("renewable energy", "report"),
            "Renewable_Energy_Report",
        )
        self.assertEqual(
            generate_document_filename("messi and ronaldo", "comparison"),
            "Messi_And_Ronaldo_Comparison",
        )
        self.assertEqual(
            generate_document_filename("machine learning", "document"),
            "Machine_Learning_Overview",
        )

    def test_create_and_verify_docx(self):
        from mini_kio.core.document_operator import (
            create_document, verify_docx, _parse_content_blocks,
        )
        import pathlib, tempfile
        content = (
            "Machine Learning: A Brief Overview\n\n"
            "What is Machine Learning?\n"
            "Machine learning is a branch of AI where systems learn from data.\n\n"
            "Key Concepts\n- Training data\n- Models\n- Generalization\n\n"
            "Why It Matters\nML powers modern software.\n"
        )
        blocks = _parse_content_blocks(content)
        self.assertTrue(any(k == "heading" for k, _ in blocks))
        self.assertTrue(any(k == "bullet" for k, _ in blocks))
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_document("machine learning", content, artifact="report", out_dir=tmp)
        self.assertTrue(r["success"], r)
        self.assertTrue(r["filename"].startswith("Machine_Learning_Report"))
        path = pathlib.Path(r["path"])
        self.assertTrue(path.exists())
        facts = verify_docx(path)
        self.assertTrue(facts["valid_zip"])
        self.assertGreater(facts["word_count"], 10)

    def test_collision_never_overwrites(self):
        from mini_kio.core.document_operator import (
            create_document, _ensure_unique,
        )
        import pathlib, tempfile
        tmp = pathlib.Path(tempfile.mkdtemp())
        p1 = _ensure_unique(tmp / "Doc.docx")
        p1.write_bytes(b"x")
        p2 = _ensure_unique(tmp / "Doc.docx")
        self.assertNotEqual(p1, p2)
        self.assertEqual(p2.name, "Doc (2).docx")

    def test_create_document_routing(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify(
            "create a word document about machine learning",
            "create a word document about machine learning",
        )
        self.assertEqual(d.action, "create_document")
        self.assertEqual(d.target, "machine learning")
        self.assertEqual((d.metadata or {}).get("artifact"), "document")
        d2 = p._classifier.classify(
            "make a word file comparing messi and ronaldo",
            "make a word file comparing messi and ronaldo",
        )
        self.assertEqual(d2.action, "create_document")
        self.assertEqual((d2.metadata or {}).get("artifact"), "comparison")
        d3 = p._classifier.classify(
            "create a report on renewable energy and make it concise",
            "create a report on renewable energy and make it concise",
        )
        self.assertEqual(d3.action, "create_document")
        self.assertEqual(d3.target, "renewable energy")
        self.assertEqual((d3.metadata or {}).get("style"), "concise")

    def test_no_markdown_leak_in_document_xml(self):
        from mini_kio.core.document_operator import _sanitize_markdown
        dirty = "## Headers\n**bold** text with `code` and ```fences```\n"
        clean = _sanitize_markdown(dirty)
        self.assertNotIn("##", clean)
        self.assertNotIn("**", clean)
        self.assertNotIn("`", clean)
        self.assertNotIn("```", clean)



class ReferentStyleModifierTest(unittest.TestCase):
    """Live-found (document-creation revalidation): 'create a word document
    about renewable energy and make IT concise' spliced the active entity into
    the style modifier -> 'make notepad concise' -> filename
    Renewable_Energy_And_Make_Notepad_Concise_Overview.docx. A pronoun in a
    QUALITY-MODIFIER construction is a style/format modifier, never a target
    referent. Real referents ("close it") must still resolve.
    """

    def _ctx_with_entity(self, entity):
        from mini_kio.core.context_manager import SessionContext
        cm = SessionContext(session_id="guard-test")
        cm.update({"success": True, "action": "open_app", "target": entity},
                  f"open {entity}")
        return cm

    def test_style_modifier_it_is_not_resolved(self):
        cm = self._ctx_with_entity("notepad")
        resolved = cm.resolved_text(
            "create a word document about renewable energy and make it concise"
        )
        self.assertIn("make it concise", resolved)
        self.assertNotIn("make notepad", resolved)

    def test_style_modifier_variants(self):
        cm = self._ctx_with_entity("word")
        for q in (
            "keep it short",
            "turn it professional",
            "write it as a poem",
            "make it a list",
        ):
            self.assertEqual(cm.resolved_text(q), q, q)

    def test_real_referent_still_resolves(self):
        cm = self._ctx_with_entity("notepad")
        self.assertEqual(cm.resolved_text("close it"), "close notepad")
        self.assertEqual(cm.resolved_text("focus it"), "focus notepad")


class SemanticEditVariantsTest(unittest.TestCase):
    """Natural-language semantic edit actions (Section 13 directive):
    "select everything", "copy the selected text", "save the document" must
    converge to their semantic action with the key combo in metadata — never
    a mechanical-key vocabulary leak or an LLM/conversation fallback."""

    def _classify(self, text):
        from mini_kio.core.pipeline import _IntentClassifier
        return _IntentClassifier().classify(text, text)

    def test_semantic_edit_variants_route(self):
        cases = {
            "select everything": ("select_all", "ctrl+a"),
            "select the whole thing": ("select_all", "ctrl+a"),
            "select all text": ("select_all", "ctrl+a"),
            "copy the selected text": ("copy", "ctrl+c"),
            "copy everything": ("copy", "ctrl+c"),
            "save the document": ("save", "ctrl+s"),
            "save the file": ("save", "ctrl+s"),
            "save all": ("save", "ctrl+s"),
            "paste it here": ("paste", "ctrl+v"),
            "undo that": ("undo", "ctrl+z"),
            "undo the last action": ("undo", "ctrl+z"),
            "redo this": ("redo", "ctrl+y"),
        }
        for phrase, (action, combo) in cases.items():
            d = self._classify(phrase)
            self.assertEqual(d.intent_type.value, "desktop_action", phrase)
            self.assertEqual(d.action, action, phrase)
            self.assertEqual(d.metadata.get("combo"), combo, phrase)

    def test_pronoun_edit_actions_survive_context_resolution(self):
        # Live-found regression: context resolution spliced "copy it" into
        # "copy notepad", destroying the semantic action. Edit-action verbs
        # must keep their pronoun so the classifier recognizes the action and
        # the executor resolves the target from context.
        from mini_kio.core.context_manager import get_session_context
        from mini_kio.core.pipeline import _IntentClassifier
        ctx = get_session_context("test-edit-pronoun")
        ctx.active_entity = "notepad"
        ctx.last_target = "notepad"
        c = _IntentClassifier()
        for phrase, action in (
            ("copy it", "copy"), ("paste it", "paste"), ("save it", "save"),
            ("undo it", "undo"), ("redo it", "redo"),
        ):
            resolved = ctx.resolved_text(phrase)
            d = c.classify(resolved, phrase)
            self.assertEqual(d.intent_type.value, "desktop_action", (phrase, resolved))
            self.assertEqual(d.action, action, (phrase, resolved))

    def test_generic_document_target_routes_to_create_document(self):
        # "do a comparison of A and B and put it in a new document" (no app
        # named) is a DOCUMENT-CREATION request — the document operator writes
        # a real artifact, never keystroke-TYPE into an app called "document".
        from mini_kio.core.pipeline import _IntentClassifier
        c = _IntentClassifier()
        d = c.classify(
            "do a small comparison of messi and ronaldo and put it in a new document",
            "do a small comparison of messi and ronaldo and put it in a new document",
        )
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "create_document")
        self.assertEqual(d.metadata.get("artifact"), "comparison")
        self.assertEqual(d.metadata.get("subject"), "messi and ronaldo")

    def test_named_app_document_stays_type_generate(self):
        # Same content request but with a NAMED app keeps the TYPE-into-app
        # path (generate the content, then type it into the fresh Notepad).
        from mini_kio.core.pipeline import _IntentClassifier
        c = _IntentClassifier()
        d = c.classify(
            "do a small comparison on KTU 2024 scheme and KTU 2019 scheme and type it into notepad",
            "do a small comparison on KTU 2024 scheme and KTU 2019 scheme and type it into notepad",
        )
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "type")
        self.assertEqual(d.target, "notepad")
        self.assertTrue(d.metadata.get("generate"))

    def test_literal_type_never_generates(self):
        from mini_kio.core.pipeline import _IntentClassifier
        c = _IntentClassifier()
        d = c.classify("type hello into notepad", "type hello into notepad")
        self.assertEqual(d.action, "type")
        self.assertFalse(d.metadata.get("generate"))

    def test_research_facts_kio_self_short_circuit(self):
        # KIO-self content requests must NEVER reach the web research chain.
        from mini_kio.core.pipeline import _ExecutionCoordinator
        coord = _ExecutionCoordinator()
        self.assertEqual(coord._research_facts("write a document about KIO capabilities"), "")
        self.assertEqual(coord._research_facts("what is kio's architecture"), "")

    def test_conversation_reply_never_poisons_context_referent(self):
        # Live-found: a companion reply stored the USER'S OWN SENTENCE as
        # active_entity, which the next request's pronoun splice injected
        # verbatim ("put it in a new document" -> "put <whole previous
        # question> in a new document"), dragging a fresh document request
        # into the conversation path. Sentence-length referents must never be
        # stored or spliced.
        from mini_kio.core.context_manager import get_session_context
        ctx = get_session_context("test-poison-referent")
        # Simulate a conversation result whose target is the raw user sentence.
        ctx.update(
            {
                "success": True,
                "message": "I'd choose Spotify...",
                "target": "if you had to pick, would you rather recommend Spotify or YouTube Music for discovering new music",
                "action": "converse",
            },
            "if you had to pick, would you rather recommend Spotify or YouTube Music for discovering new music",
        )
        self.assertIsNone(ctx.active_entity)
        self.assertIsNone(ctx.last_target)
        # A fresh document request must survive context resolution intact.
        resolved = ctx.resolved_text(
            "do a small comparison of Messi and Ronaldo and put it in a new document"
        )
        self.assertEqual(resolved, "do a small comparison of Messi and Ronaldo and put it in a new document")
        # Concise entity referents still splice normally.
        ctx2 = get_session_context("test-poison-referent-ok")
        ctx2.update({"success": True, "message": "Done.", "target": "notepad", "action": "open_app"}, "open notepad")
        self.assertEqual(ctx2.active_entity, "notepad")
        self.assertEqual(ctx2.resolved_text("close it"), "close notepad")

    def test_research_facts_grounds_non_kio_query(self):
        # A factual non-KIO query goes through the router; a provider miss
        # must degrade to "" (never raise, never block content generation).
        from mini_kio.core.pipeline import _ExecutionCoordinator
        coord = _ExecutionCoordinator()
        facts = coord._research_facts("compare KTU 2024 and 2019 schemes")
        self.assertIsInstance(facts, str)

    def test_semantic_edit_with_explicit_app_target(self):
        d = self._classify("save the file in word")
        self.assertEqual(d.intent_type.value, "desktop_action")
        self.assertEqual(d.action, "save")
        self.assertEqual(d.metadata.get("combo"), "ctrl+s")

    def test_knowledge_shapes_stay_out(self):
        for phrase in (
            "copy files to usb", "save money", "what is copywriting",
            "how to save a file", "select all files", "select all files in downloads",
        ):
            d = self._classify(phrase)
            self.assertNotEqual(d.intent_type.value, "desktop_action", phrase)
            self.assertNotEqual(d.action, "copy", phrase)


class CameraCapabilityTest(unittest.TestCase):
    """Camera capability: semantic family + generic UWP native resolution.

    - "open camera" / "launch my camera" must resolve the NATIVE installed
      camera app (UWP discovery), never a .com website.
    - "take a picture" / "capture a photo" must route to the camera capture
      capability, never to the LLM/web.
    - The executor must claim capture success ONLY when a photo file actually
      appeared (truthful verification), otherwise report the limitation.
    """

    def _classify(self, text):
        from mini_kio.core.pipeline import _IntentClassifier
        return _IntentClassifier().classify(text, text)

    def test_capture_phrasing_routes_to_camera_capability(self):
        for q in (
            "take a picture",
            "take a photo with the camera",
            "capture a photo",
            "snap a picture",
            "take a selfie",
            "take a picture with the camera",
        ):
            d = self._classify(q)
            self.assertEqual(d.intent_type.value, "desktop_action", q)
            self.assertEqual(d.action, "camera", q)
            self.assertEqual(d.target, "capture", q)
            self.assertEqual(d.metadata.get("camera_action"), "capture", q)

    def test_open_camera_routes_native_app(self):
        for q in ("open the camera", "open camera", "launch my camera"):
            d = self._classify(q)
            # Native open path (target normalized to "camera"); the executor's
            # native-first discovery resolves the UWP app — never a website.
            self.assertEqual(d.intent_type.value, "desktop_open", q)
            self.assertEqual(d.target, "camera", q)

    def test_what_is_a_camera_stays_knowledge(self):
        d = self._classify("what is a camera")
        self.assertNotEqual(d.intent_type.value, "desktop_action")
        self.assertNotEqual(d.action, "camera")

    def _reset_uwp_cache(self):
        from mini_kio.core.app_operator import _UWP_APPS_CACHE
        _UWP_APPS_CACHE["ts"] = 0.0
        _UWP_APPS_CACHE["failed_ts"] = 0.0
        _UWP_APPS_CACHE["apps"] = []

    @mock.patch("mini_kio.core.execution_boundary.execute_action")
    @mock.patch("mini_kio.core.app_operator.subprocess.run")
    def test_uwp_discovery_finds_native_camera(self, mock_subprocess_run, mock_exec):
        # Get-StartApps returns the real Windows Camera AUMID.
        mock_subprocess_run.return_value = mock.Mock(
            stdout="Camera|Microsoft.WindowsCamera_8wekyb3d8bbwe!App\r\n"
                   "Photos|Microsoft.Windows.Photos_8wekyb3d8bbwe!App\r\n",
            stderr="",
        )
        self._reset_uwp_cache()
        from mini_kio.core.app_operator import _find_installed_app
        found = _find_installed_app("camera")
        self.assertIsNotNone(found)
        self.assertEqual(found["kind"], "uwp")
        self.assertIn("WindowsCamera", found["target"])
        # Routing must prefer the native route over the .com website.
        from mini_kio.core.routing_utils import get_browser_routing
        route = get_browser_routing("camera")
        self.assertEqual(route["route_type"], "native")
        self._reset_uwp_cache()

    @mock.patch("mini_kio.core.app_operator.subprocess.run")
    def test_uwp_short_fragment_never_overmatches(self, mock_subprocess_run):
        mock_subprocess_run.return_value = mock.Mock(
            stdout="Camera|Microsoft.WindowsCamera_8wekyb3d8bbwe!App\r\n"
                   "Photos|Microsoft.Windows.Photos_8wekyb3d8bbwe!App\r\n",
            stderr="",
        )
        self._reset_uwp_cache()
        from mini_kio.core.app_operator import _find_installed_app
        self.assertIsNone(_find_installed_app("a"))  # len < 3
        self.assertIsNone(_find_installed_app("zq"))  # len < 3
        self._reset_uwp_cache()

    def test_capture_executor_claims_only_verified_photo(self):
        # The camera executor must return an honest limitation when no photo
        # file appears, never a fake "Took a photo". The REAL executor is
        # exercised with all external dependencies mocked at their import
        # sites (the method imports at call time).
        from mini_kio.core.pipeline import _ExecutionCoordinator
        coord = _ExecutionCoordinator()
        dp = mock.Mock()
        dp.execute.return_value = {"success": True}
        with mock.patch(
            "mini_kio.core.app_operator._find_installed_app",
            return_value={"kind": "uwp", "target": "CameraAUMID", "display": "Camera"},
        ), mock.patch(
            "mini_kio.core.app_operator._launch_discovered",
            return_value={"success": True, "message": "Opened camera"},
        ), mock.patch("mini_kio.core.pipeline.time.sleep"), mock.patch(
            "pathlib.Path.home",
            return_value=__import__("pathlib").Path("."),
        ):
            # A non-existent Camera Roll dir → no new photos → truthful
            # limitation, never a fabricated capture.
            result = coord._exec_camera({"camera_action": "capture"}, dp)
        self.assertFalse(result["success"])
        self.assertIn("couldn't confirm", result["message"])
        self.assertTrue(result.get("camera_open"))
        self.assertNotIn("Took a photo", result["message"])


class GenericArtifactCreationTest(unittest.TestCase):
    """Generic CONTENT + ARTIFACT capability: the same semantic task model
    resolves essay/report -> .docx (Word), spreadsheet/budget -> .xlsx
    (Excel), presentation/slides -> .pptx (PowerPoint) through ONE canonical
    artifact operator — never per-app branches. Destination applications in
    the user's phrasing are FORMAT hints, and every artifact is verified."""

    def _classify(self, text):
        from mini_kio.core.pipeline import _IntentClassifier
        return _IntentClassifier().classify(text, text)

    def test_spreadsheet_and_presentation_routing(self):
        cases = [
            ("make a spreadsheet comparing X and Y in Excel", "spreadsheet"),
            ("create a presentation about black holes", "presentation"),
            ("prepare slides about KIO", "presentation"),
            ("create a budget spreadsheet", "spreadsheet"),
            ("make an excel sheet of sales data", "spreadsheet"),
            ("make a comparison of Messi and Ronaldo in Excel", "spreadsheet"),
            ("create a study guide about python", "study guide"),
            ("write an essay about climate change in Word", "essay"),
            ("create a report on the KTU 2024 scheme", "report"),
        ]
        for text, artifact in cases:
            d = self._classify(text)
            self.assertEqual(d.action, "create_document", text)
            self.assertEqual((d.metadata or {}).get("artifact"), artifact, text)

    def test_destination_app_tail_is_stripped_from_subject(self):
        d = self._classify("make a spreadsheet comparing X and Y in Excel")
        self.assertNotIn("excel", d.target.lower())
        self.assertEqual(d.target.strip().lower(), "x and y")

    def test_bare_artifact_asks_for_topic(self):
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        d = p._classifier.classify("create a study guide", "create a study guide")
        self.assertEqual(d.action, "create_document")
        self.assertEqual((d.metadata or {}).get("artifact"), "study guide")
        self.assertEqual((d.metadata or {}).get("subject"), "")

    def test_pptx_per_slide_notes_and_brace_safe_content(self):
        # Capability-quality: speaker notes are native to the presentation
        # format; each slide must reference ITS OWN notes slide (never a
        # shared notesSlide1), and speaker-notes text containing literal { }
        # braces must not crash the builder (.format() would misread them).
        import pathlib, tempfile, zipfile
        from mini_kio.core.artifact_operator import create_artifact, verify_pptx
        content = (
            "Intro\n- Uses {x} and {y} syntax\n- Second point\n"
            "Why It Matters!\n- Point A\n- Point B\n"
            "Conclusion\n- Wrap up\n"
        )
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact("brace test", content, artifact="presentation", out_dir=tmp)
        self.assertTrue(r.get("success"), r.get("message"))
        f = pathlib.Path(r["path"])
        facts = verify_pptx(f)
        # Rich decks add a title slide in front of the 3 content slides.
        self.assertEqual(facts.get("slide_count"), 4)
        self.assertEqual(facts.get("notes_count"), 4)
        with zipfile.ZipFile(str(f)) as z:
            rel1 = z.read("ppt/slides/_rels/slide2.xml.rels").decode("utf-8")
            self.assertIn("notesSlide2.xml", rel1)  # per-slide, not shared
            # The brace content lives on content slide 2 (slide 1 = title).
            n2 = z.read("ppt/notesSlides/notesSlide2.xml").decode("utf-8")
            self.assertIn("{x}", n2)  # braces preserved, no crash

    def test_xlsx_meaningful_sheet_name_and_header_style(self):
        # Capability-quality: a spreadsheet is a REAL structured artifact —
        # meaningful sheet name derived from the subject (no "Budget_Budget"
        # duplication) and a bold header row (style s=1), not a text dump.
        import pathlib, tempfile, zipfile
        from mini_kio.core.artifact_operator import (
            create_artifact, _sheet_name_from_subject,
        )
        self.assertEqual(_sheet_name_from_subject("budget", "budget"), "Budget")
        self.assertEqual(
            _sheet_name_from_subject("messi and ronaldo", "comparison"),
            "Messi_And_Ronaldo",
        )
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact(
            "messi vs ronaldo",
            "Player\tGoals\tAssists\nMessi\t700\t300\nRonaldo\t800\t250",
            artifact="spreadsheet",
            out_dir=tmp,
        )
        self.assertTrue(r.get("success"), r.get("message"))
        f = pathlib.Path(r["path"])
        with zipfile.ZipFile(str(f)) as z:
            wb = z.read("xl/workbook.xml").decode("utf-8")
            sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertIn('name="Messi_Vs_Ronaldo"', wb)
        header_row = sheet.split('<row r="1">')[1].split("</row>")[0]
        self.assertIn('s="1"', header_row)  # bold header style applied

    def test_code_workflow_natural_forms(self):
        # Capability-quality (live directive): "create a small Python CLI
        # calculator" and "make a simple HTML page" must route to the code
        # artifact workflow (language inferred, file opened in the editor),
        # never to conversation or to Notepad typing.
        cases = [
            ("create a small Python CLI calculator", "calculator", "python"),
            ("make a simple HTML page", "html page", "html"),
            ("create a Python project with a README", "with a readme", "python"),
            ("write some code for this in VS Code", "", "python"),
        ]
        for text, expect_subj, expect_lang in cases:
            d = self._classify(text)
            self.assertEqual(d.action, "create_document", text)
            md = d.metadata or {}
            self.assertEqual(md.get("artifact"), "code", text)
            self.assertEqual(md.get("language"), expect_lang, text)
            if expect_subj:
                self.assertIn(expect_subj, (md.get("subject") or ""), text)

    def test_code_workflow_does_not_swallow_artifacts_or_phrases(self):
        # The code-noun guard must stay BOUNDED: spreadsheets stay
        # spreadsheets, and non-code phrases ("make a paper airplane") stay
        # conversational — the code regex requires a real code noun.
        for text, action in [
            ("make a small spreadsheet comparing tea and coffee", "create_document"),
            ("make a paper airplane", "converse"),
            ("create a mess", "converse"),
        ]:
            d = self._classify(text)
            if action == "converse":
                self.assertEqual(d.action, "converse", text)
            else:
                self.assertEqual(d.action, "create_document", text)
                self.assertEqual((d.metadata or {}).get("artifact"), "spreadsheet", text)

    def test_xlsx_rich_workbook_has_formula_and_chart(self):
        # Capability-quality: a real workbook contains a SUM totals row and a
        # chart part — not a bare text grid. openpyxl writes both. (Skipped
        # when the interpreter lacks openpyxl — the stdlib fallback still
        # produces a valid workbook without charts.)
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed in this interpreter")
        import pathlib, tempfile, zipfile
        from mini_kio.core.artifact_operator import create_artifact
        content = "Category\tAmount\nRent\t1200\nFood\t400\nTransport\t150\nSavings\t300"
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact("monthly budget", content, artifact="spreadsheet", out_dir=tmp)
        self.assertTrue(r.get("success"), r.get("message"))
        f = pathlib.Path(r["path"])
        with zipfile.ZipFile(str(f)) as z:
            names = set(z.namelist())
            sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8", errors="replace")
        self.assertTrue(any("chart" in n for n in names), "chart part missing")
        self.assertIn("SUM", sheet.upper())

    def test_docx_rich_has_table_footer_and_headings(self):
        # Capability-quality: a Word document has real structure — a markdown
        # table becomes a real Word table, footer carries a page-number field,
        # and short heading lines render as heading paragraphs. (Skipped when
        # the interpreter lacks python-docx; the stdlib fallback still works.)
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed in this interpreter")
        import pathlib, tempfile
        from mini_kio.core.artifact_operator import create_artifact
        content = (
            "Introduction\nRenewable energy is growing fast.\n\n"
            "Key Sources\n- Solar\n- Wind\n\n"
            "| Source | Share |\n| Solar | 30% |\n| Wind | 25% |\n\n"
            "Conclusion\nThe transition is accelerating.\n"
        )
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact("renewable energy report", content, artifact="report", out_dir=tmp)
        self.assertTrue(r.get("success"), r.get("message"))
        doc = Document(str(r["path"]))
        self.assertGreaterEqual(len(doc.tables), 1, "no Word table")
        footer = doc.sections[0].footer.paragraphs[0].text if doc.sections[0].footer.paragraphs else ""
        self.assertIn("Page", footer)
        heading_texts = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        self.assertTrue(any("Introduction" in t or "Conclusion" in t for t in heading_texts))

    def test_pptx_rich_deck_has_slides_notes_and_numbers(self):
        # Capability-quality: python-pptx decks have per-slide speaker notes,
        # explicit SLIDE: marker support, and slide-number textboxes.
        try:
            import pptx  # noqa: F401
        except ImportError:
            self.skipTest("python-pptx not installed in this interpreter")
        import pathlib, tempfile
        from mini_kio.core.artifact_operator import create_artifact, verify_pptx
        content = (
            "SLIDE: Overview\n- Topic intro\n- Why it matters\n\n"
            "SLIDE: Comparison\n| A | B |\n| 1 | 2 |\n\n"
            "SLIDE: Conclusion\n- Wrap up\n"
        )
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact("ai overview", content, artifact="presentation", out_dir=tmp)
        self.assertTrue(r.get("success"), r.get("message"))
        facts = verify_pptx(pathlib.Path(r["path"]))
        # Rich decks add a title slide in front of the 3 content slides.
        self.assertEqual(facts.get("slide_count"), 4)
        self.assertEqual(facts.get("notes_count"), 4)

    def test_non_artifact_phrases_stay_conversational(self):
        for q in ("make a paper airplane", "create a mess", "make a move"):
            d = self._classify(q)
            self.assertNotEqual(d.action, "create_document", q)

    def test_bare_artifact_subject_not_leaked_from_phantom_trailing_noun(self):
        # Live-found regression: "create a study guide" was split as a phantom
        # trailing noun (subject="study"/noun="guide") and the leaked subject
        # survived the guard, so the bare-artifact route reported subject=
        # "study" instead of "". A bare artifact has NO topic — the executor
        # must ask. Same guard applies to other adjectives that are NOT topics.
        from mini_kio.core.pipeline import Pipeline
        p = Pipeline()
        for q in ("create a study guide", "make a quick comparison",
                  "write a short poem", "create a detailed report"):
            d = p._classifier.classify(q, q)
            self.assertEqual(d.action, "create_document", q)
            md = d.metadata or {}
            # "quick"/"short"/"detailed" are styles, never topics; a bare
            # artifact with no topic marker keeps subject empty.
            self.assertNotIn(
                md.get("subject"), ("study", "quick", "short", "detailed"), q
            )
            self.assertEqual(md.get("subject"), "", q)

    def test_create_document_response_is_natural_not_execution_log(self):
        # Capability-quality (live directive): confirmations must read like
        # KIO, not like a build log. The executor keeps the truthful facts but
        # wraps them in plain language ("Done — I made you a fresh spreadsheet
        # and opened it."). No bare "Created X — 3 rows." log string.
        from unittest import mock
        from mini_kio.core.pipeline import _ExecutionCoordinator
        from mini_kio.core.pipeline.types import IntentType, RoutingDecision
        coord = _ExecutionCoordinator()
        decision = RoutingDecision(
            intent_type=IntentType.DESKTOP_ACTION,
            action="create_document",
            target="tea and coffee",
            raw_text="make a small spreadsheet comparing tea and coffee in Excel",
            normalized_text="make a small spreadsheet comparing tea and coffee in Excel",
            metadata={"artifact": "spreadsheet", "subject": "tea and coffee", "style": ""},
        )
        fake_result = {
            "success": True,
            "filename": "Tea_And_Coffee_Spreadsheet.xlsx",
            "row_count": 3,
            "cell_count": 12,
            "path": "C:/fakepath/Tea_And_Coffee_Spreadsheet.xlsx",
            "artifact": "spreadsheet",
            "subject": "tea and coffee",
        }
        with mock.patch.object(
            coord, "_generate_content", return_value="X\tY\n1\t2\n3\t4\n"), \
                mock.patch(
                    "mini_kio.core.artifact_operator.create_artifact",
                    return_value=fake_result,
                ), mock.patch(
                    "mini_kio.core.artifact_operator.open_artifact", return_value=True
                ):
            out = coord.execute(
                "desktop_action",
                {
                    "action": "create_document",
                    "target": "tea and coffee",
                    "metadata": {"artifact": "spreadsheet", "subject": "tea and coffee"},
                },
                decision,
            )
        msg = out.get("message", "")
        self.assertIn("spreadsheet", msg.lower())
        self.assertNotRegex(msg, r"^Created \S+ — \d+ (row|word|slide|line)s?\.$")

    def test_save_as_clause_is_single_document_intent(self):
        # Live-found: "write a short essay about X and save it as a Word
        # document" was misrouted to multi_step/TYPE (both "write" and "save"
        # look like verbs). It is ONE document-creation intent — the artifact
        # operator persists the file itself. The save-as clause maps to the
        # artifact FORMAT, and genuine multi-step commands stay multi-step.
        cases = [
            ("write a short essay about renewable energy and save it as a Word document", "essay"),
            ("write an essay about climate change and save it as a word doc", "essay"),
            ("draft a report about KIO and save it as a document", "report"),
            ("create a presentation about AI and save it as a powerpoint", "presentation"),
            ("make a spreadsheet about expenses and save it as an excel file", "spreadsheet"),
        ]
        for text, artifact in cases:
            d = self._classify(text)
            self.assertEqual(d.action, "create_document", text)
            self.assertEqual((d.metadata or {}).get("artifact"), artifact, text)
            self.assertNotIn("save", d.target.lower(), text)
        # Genuine multi-step must NOT be swallowed by the save-as detector.
        for text in ("open chrome and search for cats and dogs", "open github and youtube"):
            d = self._classify(text)
            self.assertEqual(d.action, "multi_step", text)

    def test_create_xlsx_spreadsheet(self):
        from mini_kio.core.artifact_operator import create_artifact, verify_xlsx
        import pathlib, tempfile
        content = "Category\tAmount\nRent\t1200\nFood\t400\nTransport\t150"
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact("monthly budget", content, artifact="spreadsheet", out_dir=tmp)
        self.assertTrue(r["success"], r)
        self.assertTrue(r["filename"].endswith(".xlsx"), r)
        self.assertIn("Monthly_Budget", r["filename"])
        path = pathlib.Path(r["path"])
        self.assertTrue(path.exists())
        facts = verify_xlsx(path)
        self.assertTrue(facts["valid_zip"])
        self.assertGreaterEqual(facts["row_count"], 4)  # header + 3 data rows
        self.assertGreaterEqual(facts["cell_count"], 8)

    def test_create_pptx_presentation(self):
        from mini_kio.core.artifact_operator import create_artifact, verify_pptx
        import pathlib, tempfile
        content = (
            "Introduction\n- What is a black hole\n- How they form\n"
            "Key Concepts\n- Event horizon\n- Singularity\n"
            "Why It Matters\n- Spacetime\n- Cosmology"
        )
        tmp = pathlib.Path(tempfile.mkdtemp())
        r = create_artifact("black holes", content, artifact="presentation", out_dir=tmp)
        self.assertTrue(r["success"], r)
        self.assertTrue(r["filename"].endswith(".pptx"), r)
        self.assertIn("Black_Holes", r["filename"])
        path = pathlib.Path(r["path"])
        self.assertTrue(path.exists())
        facts = verify_pptx(path)
        self.assertTrue(facts["valid_zip"])
        self.assertGreaterEqual(facts["slide_count"], 3)

    def test_artifact_files_are_wellformed_ooxml(self):
        # Every XML/rels part inside the produced artifacts must be well-formed
        # (this is what stops Word/Excel/PowerPoint repair prompts).
        from mini_kio.core.artifact_operator import create_artifact
        import pathlib, tempfile, zipfile, xml.dom.minidom
        tmp = pathlib.Path(tempfile.mkdtemp())
        create_artifact("b", "Category\tAmount\nRent\t1", artifact="spreadsheet", out_dir=tmp)
        create_artifact("s", "Intro\n- one\n- two", artifact="presentation", out_dir=tmp)
        create_artifact("d", "Title\nBody text.", artifact="document", out_dir=tmp)
        for f in tmp.iterdir():
            with zipfile.ZipFile(str(f)) as zf:
                self.assertIsNone(zf.testzip())
                for name in zf.namelist():
                    if name.endswith((".xml", ".rels")):
                        xml.dom.minidom.parseString(zf.read(name))

    def test_pre_noun_modifier_variants_route_create_document(self):
        # Live-found (post-commit smoke): "make a SMALL spreadsheet comparing
        # tea and coffee" fell through to conversation because the create
        # regexes only allowed new/fresh/another before the artifact noun.
        # Any bounded pre-noun modifier must keep artifact routing intact.
        cases = [
            ("make a small spreadsheet comparing tea and coffee in Excel", "spreadsheet"),
            ("make a simple spreadsheet comparing tea and coffee", "spreadsheet"),
            ("make a small spreadsheet comparing tea and coffee", "spreadsheet"),
            ("create a small spreadsheet about my expenses", "spreadsheet"),
            ("create a quick presentation about black holes", "presentation"),
            ("write a short essay about renewable energy", "essay"),
        ]
        for text, artifact in cases:
            d = self._classify(text)
            self.assertEqual(d.action, "create_document", text)
            self.assertEqual((d.metadata or {}).get("artifact"), artifact, text)

    def test_compare_form_infers_artifact_from_leading_noun(self):
        # "make a spreadsheet comparing X and Y" without an "in Excel" tail
        # names the artifact in the LEADING clause: it must build an .xlsx,
        # never a comparison .docx. A generic container ("word file") stays a
        # comparison document inside the Word container.
        d = self._classify("make a spreadsheet comparing tea and coffee")
        self.assertEqual((d.metadata or {}).get("artifact"), "spreadsheet")
        d2 = self._classify("make a presentation comparing linux and windows")
        self.assertEqual((d2.metadata or {}).get("artifact"), "presentation")
        d3 = self._classify("make a word file comparing messi and ronaldo")
        self.assertEqual((d3.metadata or {}).get("artifact"), "comparison")

    def test_creative_phrases_with_modifiers_stay_conversational(self):
        # The modifier slot must not swallow non-artifact requests.
        for text in ("make a paper airplane", "create a mess", "make a move",
                     "write a letter to my mom", "make a quick trip to the store"):
            d = self._classify(text)
            self.assertNotEqual(d.action, "create_document", text)


class CasualFragmentRoutingTest(unittest.TestCase):
    """Live-found: "Yoo" (message-initial capital) became an ENTITY_QUERY
    and returned an unrelated UFC fighter biography. Message-initial
    capitalization is a writing convention, not proper-noun evidence —
    casual/greeting/social fragments must stay on the conversation family
    while genuine entities (Messi, Spotify, ChatGPT) keep the entity path."""

    def _classify(self, text):
        from mini_kio.core.pipeline import _IntentClassifier
        return _IntentClassifier().classify(text, text)

    def test_casual_fragments_never_become_entity_queries(self):
        # The exact live regression plus a broad casual family, all typed with
        # a message-initial capital letter (the failing condition).
        casual = (
            "Yoo", "Yooo", "Yo", "Hey", "Heyy", "Hi", "Hello", "Sup",
            "Lol", "Lmao", "Wow", "Damn", "Cool", "Nice", "Okay", "K",
            "Alright", "Thanks", "Ty", "Bro", "Dude", "Bruh", "Yikes",
            "Hmm", "Oh", "Ah", "Yeah", "Yep", "Nope", "Nah", "Omg",
            # Punctuation variants are the most natural greeting forms — they
            # must not leak through the guard ("yoo!" != "yoo" raw).
            "Yoo!", "Yoo?", "Hey!", "Hi!", "Lol!", "Wow!", "Oh wow", "Haha yeah",
        )
        for q in casual:
            d = self._classify(q)
            self.assertIn(
                d.intent_type.value, ("conversation", "greeting", "social"),
                f"{q!r} must stay conversational, got {d.intent_type.value}",
            )
            self.assertNotEqual(
                d.intent_type.value, "entity_query", f"{q!r} became an entity query",
            )

    def test_genuine_entities_still_route_to_entity_query(self):
        for q in ("Messi", "Ronaldo", "Spotify", "ChatGPT", "One hundred years of solitude"):
            d = self._classify(q)
            self.assertEqual(d.intent_type.value, "entity_query", q)

    def test_entities_starting_with_casual_word_keep_entity_path(self):
        # "Nice France", "Good Charlotte", "Great Barrier Reef": the all-casual
        # guard fires ONLY when every word is a casual fragment, so real
        # entities that merely START with a casual token keep the entity path.
        for q in ("Nice France", "Good Charlotte", "Great Barrier Reef", "Cool Runnings"):
            d = self._classify(q)
            self.assertEqual(d.intent_type.value, "entity_query", q)

    def test_full_pipeline_yoo_is_conversation(self):
        from mini_kio.core.pipeline import Pipeline
        res = Pipeline().run("Yoo", session_id="test-casual-yoo")
        self.assertEqual(res.get("action"), "converse")
        self.assertNotIn("UFC", str(res.get("message", "")))

    def test_repeated_vowel_casual_forms_stay_conversational(self):
        # "Yoooooo" / "Heeeey" are emphasis-stretched casual fragments, not
        # proper nouns. The entity heuristic must fold runs of 3+ repeats
        # before the capitalization check, while keeping double-letter
        # entities ("Messi", "Ronaldo") intact.
        for q in ("Yoooooo", "Yooo", "Heeeey", "Wooow", "Daaaamn", "Lool", "Sssup"):
            d = self._classify(q)
            self.assertIn(
                d.intent_type.value, ("conversation", "greeting", "social"),
                f"{q!r} must stay conversational, got {d.intent_type.value}",
            )
        # Double-letter genuine entities must NOT be collapsed away.
        for q in ("Messi", "Ronaldo", "Spotify"):
            d = self._classify(q)
            self.assertEqual(d.intent_type.value, "entity_query", q)

    def test_conversation_reply_never_becomes_next_referent(self):
        # Live-found class: a conversational reply's target is the user's own
        # sentence. Even a SHORT one-word reply ("sure") must never become the
        # referent spliced into the next command ("open it" -> "open sure").
        # Gated by action KIND, not length: conversation/identity/operational
        # results never feed the referent store.
        from mini_kio.core.context_manager import get_session_context
        ctx = get_session_context("test-ref-gate")
        ctx.active_entity = None
        ctx.last_target = None
        ctx.update({"success": True, "action": "converse", "target": "sure"}, "sure")
        self.assertIsNone(ctx.active_entity)
        self.assertIsNone(ctx.last_target)
        ctx.update({"success": True, "action": "identity", "target": "KIO"}, "what is kio")
        self.assertIsNone(ctx.active_entity)
        # Real executed actions still become referents.
        ctx.update({"success": True, "action": "open_app", "target": "notepad"}, "open notepad")
        self.assertEqual(ctx.active_entity, "notepad")
        ctx.update(
            {"success": True, "action": "create_document", "target": "Messi_Comparison.docx"},
            "create doc",
        )
        self.assertIn("messi", ctx.active_entity)
        self.assertIn("comparison", ctx.resolved_text("close it"))



class BrowserModalityGatingTest(unittest.TestCase):
    """Live-found (2026-08-12 revalidation): "open youtube in comet" opened a
    Chrome tab AND launched Comet — the Chrome extension connector served an
    explicit non-default-browser request. The connector may only serve the
    DEFAULT browser; explicit edge/comet/firefox/brave requests must go
    straight to their own browser binary (never a silent Chrome substitution).
    """

    def _make_conn(self):
        from types import SimpleNamespace
        conn = mock.Mock()
        conn.is_connected.return_value = True
        conn.open_tab = mock.AsyncMock(return_value=SimpleNamespace(
            success=True, tab_id="t1", url="https://youtube.com",
            title="YouTube", window_id=1,
        ))
        return conn

    def _patch_env(self, default_browser):
        from mini_kio.core import app_operator as ao
        from mini_kio.core import routing_utils as ru
        from mini_kio.core import command_router as cr
        from mini_kio.core import config as cfg
        from unittest import mock
        return [
            mock.patch.object(ao, "_find_in_registry",
                              return_value={"exe": "comet.exe",
                                            "process": "comet.exe"}),
            mock.patch.object(ao, "_resolve_path",
                              return_value=r"C:\fake\comet.exe"),
            mock.patch.object(ao, "subprocess"),
            mock.patch.object(ao, "_refine_pid_windows", return_value=9876),
            mock.patch.object(ru, "register_browser_capability"),
            mock.patch.object(cr, "_get_connector"),
            mock.patch.object(cfg, "DEFAULT_BROWSER", default_browser),
        ]

    def test_nondefault_browser_skips_connector(self):
        # DEFAULT_BROWSER=chrome, request is comet: the Chrome connector must
        # NOT be used (no Chrome tab side effect); Comet binary is launched.
        from types import SimpleNamespace
        from mini_kio.core import app_operator as ao
        conn = self._make_conn()
        patches = self._patch_env("chrome")
        with patches[6], patches[5] as get_conn, patches[0], patches[1], \
                patches[2] as sp, patches[3], patches[4]:
            get_conn.return_value = conn
            sp.Popen.return_value = SimpleNamespace(pid=999)
            result = ao.execute_capability(
                "comet::open_url::https://youtube.com::youtube"
            )
        conn.open_tab.assert_not_called()
        self.assertIn("Comet", result["message"])
        self.assertTrue(result["success"])

    def test_default_browser_uses_connector(self):
        # DEFAULT_BROWSER=chrome, request is chrome: the connector path runs
        # (tab-level open with verification), no binary relaunch.
        from mini_kio.core import app_operator as ao
        conn = self._make_conn()
        patches = self._patch_env("chrome")
        with patches[6], patches[5] as get_conn, patches[0], patches[1], \
                patches[2] as sp, patches[3], patches[4]:
            get_conn.return_value = conn
            result = ao.execute_capability(
                "chrome::open_url::https://youtube.com::youtube"
            )
        conn.open_tab.assert_awaited_once()
        self.assertIn("YouTube", result["message"])



if __name__ == "__main__":
    unittest.main()
